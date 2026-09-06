import io
from datetime import date
from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from apps.crm.models import CrmCredit, CrmPayment, CustomerRecord, credit_balance
from apps.inventory.models import Branch
from apps.sales.models import Customer, Sale, Transaction
from apps.users.models import User


class CrmAccessTest(TestCase):
    """CRM is an admin + finance screen; everyone else is locked out."""

    def setUp(self):
        self.admin = User.objects.create_user(username='crm_admin', password='pw', role='admin')
        self.accountant = User.objects.create_user(username='crm_acct', password='pw', role='accountant')
        self.sales_rep = User.objects.create_user(username='crm_rep', password='pw', role='sales_rep')

    def test_admin_and_accountant_can_open_page(self):
        for user in (self.admin, self.accountant):
            self.client.force_login(user)
            response = self.client.get(reverse('crm:customer_records'))
            self.assertEqual(response.status_code, 200, user.role)

    def test_other_roles_are_blocked(self):
        self.client.force_login(self.sales_rep)
        self.assertEqual(self.client.get(reverse('crm:customer_records')).status_code, 403)
        self.assertEqual(self.client.get('/api/crm-records/').status_code, 403)


class CrmRecordApiTest(TestCase):
    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct2', password='pw', role='accountant')
        self.client.force_login(self.accountant)

    def test_create_list_and_summary(self):
        payload = {
            'date': '2026-08-01',
            'receipt_number': 'RC-0012',
            'efd_receipt_number': '35EFD9921',
            'customer_name': 'Kibo Traders',
            'tin': '109-882-441',
            'sales_amount': '1250000.00',
        }
        response = self.client.post('/api/crm-records/', payload, content_type='application/json')
        self.assertEqual(response.status_code, 201, response.content)

        record = CustomerRecord.objects.get()
        self.assertEqual(record.customer_name, 'Kibo Traders')
        self.assertEqual(record.created_by, self.accountant)

        listing = self.client.get('/api/crm-records/?search=Kibo').json()
        self.assertEqual(listing['count'], 1)
        self.assertEqual(listing['results'][0]['efd_receipt_number'], '35EFD9921')

        # A search that matches nothing filters the row out
        self.assertEqual(self.client.get('/api/crm-records/?search=nobody').json()['count'], 0)

        summary = self.client.get('/api/crm-records/summary/').json()
        self.assertEqual(summary['records'], 1)
        self.assertEqual(summary['customers'], 1)
        self.assertEqual(float(summary['total_amount']), 1250000.00)

    def test_csv_export_uses_filters(self):
        CustomerRecord.objects.create(date=date(2026, 8, 1), customer_name='Kibo Traders', sales_amount=100)
        CustomerRecord.objects.create(date=date(2026, 8, 5), customer_name='Mwanza Const', sales_amount=200)

        response = self.client.get(reverse('crm:export'), {'search': 'Kibo'})
        self.assertEqual(response.status_code, 200)
        body = response.content.decode()
        self.assertIn('Kibo Traders', body)
        self.assertNotIn('Mwanza Const', body)


class CrmImportFromSalesTest(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username='crm_admin2', password='pw', role='admin')
        self.client.force_login(self.admin)
        branch = Branch.objects.create(name='Main', address='Dar es Salaam')
        customer = Customer.objects.create(name='Kibo Traders', phone='0700000000')
        self.sale = Sale.objects.create(
            invoice_number='INV-CRM-1', branch=branch, customer=customer,
            status='approved', total_amount=500000,
        )

    def test_a_sale_reaches_the_register_without_being_imported(self):
        """Sales sync themselves now, so there is nothing left to pull in by hand."""
        record = CustomerRecord.objects.get()
        self.assertEqual(record.customer_name, 'Kibo Traders')
        self.assertEqual(record.receipt_number, 'INV-CRM-1')
        self.assertEqual(record.sales_amount, 500000)

        self.assertEqual(self.client.get('/api/crm-records/available_sales/').json(), [],
                         'an already-synced sale must not be offered for import')

    def test_editing_a_record_never_writes_back_to_the_sale(self):
        """The register is downstream of sales, never upstream."""
        record = CustomerRecord.objects.get()
        record.customer_name = 'Kibo Traders Ltd'
        record.sales_amount = 600000
        record.save()

        self.sale.refresh_from_db()
        self.assertEqual(self.sale.total_amount, 500000)
        self.assertEqual(self.sale.customer.name, 'Kibo Traders')

    def test_a_removed_record_can_be_pulled_back_in(self):
        """What import_sales is still for: sales that predate the sync, or a row
        somebody deleted."""
        CustomerRecord.objects.all().delete()

        available = self.client.get('/api/crm-records/available_sales/').json()
        self.assertEqual([s['invoice_number'] for s in available], ['INV-CRM-1'])

        response = self.client.post(
            '/api/crm-records/import_sales/',
            {'invoices': ['INV-CRM-1']}, content_type='application/json',
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['created'], 1)
        self.assertEqual(CustomerRecord.objects.count(), 1)

        # and never twice
        again = self.client.post(
            '/api/crm-records/import_sales/',
            {'invoices': ['INV-CRM-1']}, content_type='application/json',
        )
        self.assertEqual(again.json()['created'], 0)
        self.assertEqual(CustomerRecord.objects.count(), 1)

    def test_import_requires_invoices(self):
        response = self.client.post('/api/crm-records/import_sales/', {'invoices': []}, content_type='application/json')
        self.assertEqual(response.status_code, 400)


class CrmCustomerListTest(TestCase):
    """The CRM landing screen groups the register by customer."""

    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct4', password='pw', role='accountant')
        self.client.force_login(self.accountant)
        CustomerRecord.objects.create(date=date(2026, 8, 1), customer_name='Kibo Traders',
                                      receipt_number='RC-1', tin='', sales_amount=100)
        CustomerRecord.objects.create(date=date(2026, 8, 5), customer_name='Kibo Traders',
                                      receipt_number='RC-2', tin='109-882-441', sales_amount=250)
        CustomerRecord.objects.create(date=date(2026, 8, 3), customer_name='Mwanza Const',
                                      receipt_number='RC-3', tin='122-004-908', sales_amount=900)

    def test_customers_are_grouped_with_totals(self):
        rows = self.client.get('/api/crm-records/customers/').json()['results']
        # Most recently traded first: Kibo last bought on the 5th, Mwanza on the
        # 3rd — even though Mwanza has spent more.
        self.assertEqual([r['customer_name'] for r in rows], ['Kibo Traders', 'Mwanza Const'])

        kibo = next(r for r in rows if r['customer_name'] == 'Kibo Traders')
        self.assertEqual(kibo['records'], 2)
        self.assertEqual(float(kibo['total_amount']), 350.0)
        self.assertEqual(kibo['first_transaction'], '2026-08-01')
        self.assertEqual(kibo['last_transaction'], '2026-08-05')
        # TIN comes from the most recent record that carries one
        self.assertEqual(kibo['tin'], '109-882-441')

    def test_customer_list_respects_filters(self):
        rows = self.client.get('/api/crm-records/customers/?search=Mwanza').json()['results']
        self.assertEqual([r['customer_name'] for r in rows], ['Mwanza Const'])

        # A date window re-totals each customer rather than dropping them
        rows = self.client.get('/api/crm-records/customers/?start=2026-08-04').json()['results']
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['customer_name'], 'Kibo Traders')
        self.assertEqual(rows[0]['records'], 1)
        self.assertEqual(float(rows[0]['total_amount']), 250.0)

    def test_records_can_be_scoped_to_one_customer(self):
        rows = self.client.get('/api/crm-records/?customer=Kibo Traders').json()['results']
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r['customer_name'] == 'Kibo Traders' for r in rows))

        summary = self.client.get('/api/crm-records/summary/?customer=Kibo Traders').json()
        self.assertEqual(summary['records'], 2)
        self.assertEqual(float(summary['total_amount']), 350.0)

        # Exact match — a partial name must not leak another customer's rows
        self.assertEqual(self.client.get('/api/crm-records/?customer=Kibo').json()['results'], [])

    def test_detail_page_renders_for_a_known_customer(self):
        response = self.client.get(reverse('crm:customer_detail'), {'name': 'Kibo Traders'})
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context['exists'])
        self.assertEqual(response.context['customer_name'], 'Kibo Traders')
        self.assertEqual(response.context['tin'], '109-882-441')
        self.assertContains(response, 'Kibo Traders')

    def test_detail_page_handles_unknown_customer(self):
        response = self.client.get(reverse('crm:customer_detail'), {'name': 'Nobody Ltd'})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context['exists'])
        self.assertContains(response, 'No records found')

    def test_detail_export_is_scoped_to_the_customer(self):
        response = self.client.get(reverse('crm:export'), {'customer': 'Kibo Traders'})
        body = response.content.decode()
        self.assertIn('RC-1', body)
        self.assertIn('RC-2', body)
        self.assertNotIn('RC-3', body)

    def test_detail_page_is_closed_to_other_roles(self):
        self.client.force_login(User.objects.create_user(username='crm_rep3', password='pw', role='sales_rep'))
        response = self.client.get(reverse('crm:customer_detail'), {'name': 'Kibo Traders'})
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.get('/api/crm-records/customers/').status_code, 403)


class CrmPaymentTest(TestCase):
    """Part payments, the one-click settle, and the derived status."""

    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct7', password='pw', role='accountant')
        self.client.force_login(self.accountant)
        self.record = CustomerRecord.objects.create(
            date=date(2026, 8, 1), customer_name='Kibo Traders',
            receipt_number='RC-1', tin='109-882-441', sales_amount=1000)

    def pay(self, amount, **extra):
        payload = {'record': self.record.id, 'amount': amount,
                   'paid_on': '2026-08-02', 'method': 'cash'}
        payload.update(extra)
        return self.client.post('/api/crm-payments/', payload, content_type='application/json')

    def test_a_sale_starts_unpaid(self):
        self.assertEqual(self.record.amount_paid, Decimal('0'))
        self.assertEqual(self.record.balance, Decimal('1000'))
        self.assertEqual(self.record.payment_status, 'unpaid')

    def test_part_payment_is_kept_as_a_record(self):
        """1000 sale, 500 paid: the 500 is retained, balance is 500, part paid."""
        response = self.pay('500', reference='SLIP-77')
        self.assertEqual(response.status_code, 201, response.content)

        self.record.refresh_from_db()
        self.assertEqual(self.record.amount_paid, Decimal('500'))
        self.assertEqual(self.record.balance, Decimal('500'))
        self.assertEqual(self.record.payment_status, 'partial')

        payment = CrmPayment.objects.get()
        self.assertEqual(payment.amount, Decimal('500'))
        self.assertEqual(payment.paid_on, date(2026, 8, 2))
        self.assertEqual(payment.reference, 'SLIP-77')
        self.assertEqual(payment.created_by, self.accountant)

    def test_further_payments_accumulate_to_paid(self):
        self.pay('500')
        self.pay('300')
        self.record.refresh_from_db()
        self.assertEqual(self.record.amount_paid, Decimal('800'))
        self.assertEqual(self.record.payment_status, 'partial')

        self.pay('200')
        self.record.refresh_from_db()
        self.assertEqual(self.record.balance, Decimal('0'))
        self.assertEqual(self.record.payment_status, 'paid')
        self.assertEqual(self.record.payments.count(), 3)

    def test_payment_cannot_exceed_the_outstanding_balance(self):
        self.pay('600')
        response = self.pay('600')
        self.assertEqual(response.status_code, 400)
        self.assertIn('exceeds the outstanding balance', str(response.json()))
        self.assertEqual(CrmPayment.objects.count(), 1)

    def test_payment_must_be_positive(self):
        self.assertEqual(self.pay('0').status_code, 400)
        self.assertEqual(self.pay('-50').status_code, 400)
        self.assertEqual(CrmPayment.objects.count(), 0)

    def test_mark_paid_settles_the_remaining_balance_in_one_click(self):
        self.pay('400')
        response = self.client.post(f'/api/crm-records/{self.record.id}/mark_paid/',
                                    {}, content_type='application/json')
        self.assertEqual(response.status_code, 200, response.content)

        self.record.refresh_from_db()
        self.assertEqual(self.record.balance, Decimal('0'))
        self.assertEqual(self.record.payment_status, 'paid')
        # It records a real payment for the remainder rather than just flagging
        self.assertEqual(self.record.payments.count(), 2)
        self.assertEqual(
            sorted(p.amount for p in self.record.payments.all()),
            [Decimal('400.00'), Decimal('600.00')])

    def test_mark_paid_rejects_an_already_settled_sale(self):
        self.client.post(f'/api/crm-records/{self.record.id}/mark_paid/', {}, content_type='application/json')
        response = self.client.post(f'/api/crm-records/{self.record.id}/mark_paid/',
                                    {}, content_type='application/json')
        self.assertEqual(response.status_code, 400)
        self.assertIn('already fully paid', response.json()['detail'])
        self.assertEqual(self.record.payments.count(), 1)

    def test_removing_a_payment_reopens_the_balance(self):
        self.pay('1000')
        self.record.refresh_from_db()
        self.assertEqual(self.record.payment_status, 'paid')

        payment = CrmPayment.objects.get()
        self.assertEqual(self.client.delete(f'/api/crm-payments/{payment.id}/').status_code, 204)
        self.record.refresh_from_db()
        self.assertEqual(self.record.balance, Decimal('1000'))
        self.assertEqual(self.record.payment_status, 'unpaid')

    def test_deleting_a_sale_takes_its_payments_with_it(self):
        self.pay('500')
        self.record.delete()
        self.assertEqual(CrmPayment.objects.count(), 0)

    def test_record_list_and_history_expose_payment_state(self):
        self.pay('250')
        row = self.client.get('/api/crm-records/').json()['results'][0]
        self.assertEqual(Decimal(row['amount_paid']), Decimal('250'))
        self.assertEqual(Decimal(row['balance']), Decimal('750'))
        self.assertEqual(row['payment_status'], 'partial')

        history = self.client.get(f'/api/crm-records/{self.record.id}/payments/').json()
        self.assertEqual(len(history['payments']), 1)
        self.assertEqual(history['payment_status'], 'partial')

    def test_summary_reports_paid_and_outstanding(self):
        self.pay('400')
        summary = self.client.get('/api/crm-records/summary/').json()
        self.assertEqual(float(summary['total_amount']), 1000.0)
        self.assertEqual(float(summary['amount_paid']), 400.0)
        self.assertEqual(float(summary['balance']), 600.0)

    def test_payments_are_closed_to_other_roles(self):
        self.client.force_login(User.objects.create_user(username='crm_rep5', password='pw', role='sales_rep'))
        self.assertEqual(self.pay('100').status_code, 403)
        self.assertEqual(
            self.client.post(f'/api/crm-records/{self.record.id}/mark_paid/').status_code, 403)


class CrmCreditTest(TestCase):
    """Overpayments and deposits become credit the customer can spend later."""

    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct9', password='pw', role='accountant')
        self.client.force_login(self.accountant)
        self.sale = CustomerRecord.objects.create(
            date=date(2026, 8, 1), customer_name='Kibo Traders',
            receipt_number='RC-1', sales_amount=1000)

    def receive(self, record, amount, **extra):
        payload = {'amount': amount, 'paid_on': '2026-08-02', 'method': 'cash'}
        payload.update(extra)
        return self.client.post(f'/api/crm-records/{record.id}/receive/',
                                payload, content_type='application/json')

    def test_overpayment_settles_the_sale_and_banks_the_rest(self):
        """1,000 sale, customer hands over 1,500."""
        response = self.receive(self.sale, '1500')
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual(Decimal(str(body['applied'])), Decimal('1000'))
        self.assertEqual(Decimal(str(body['to_credit'])), Decimal('500'))

        self.sale.refresh_from_db()
        self.assertEqual(self.sale.balance, Decimal('0'))
        self.assertEqual(self.sale.payment_status, 'paid')
        # The sale's own balance never goes negative
        self.assertEqual(self.sale.amount_paid, Decimal('1000'))
        self.assertEqual(credit_balance('Kibo Traders'), Decimal('500'))

        credit = CrmCredit.objects.get()
        self.assertEqual(credit.source, 'overpayment')
        self.assertEqual(credit.customer_name, 'Kibo Traders')

    def test_credit_is_deducted_from_a_later_order(self):
        self.receive(self.sale, '1500')  # 500 left on account

        later = CustomerRecord.objects.create(
            date=date(2026, 9, 1), customer_name='Kibo Traders',
            receipt_number='RC-2', sales_amount=800)

        response = self.client.post(f'/api/crm-records/{later.id}/apply_credit/',
                                    {}, content_type='application/json')
        self.assertEqual(response.status_code, 200, response.content)

        later.refresh_from_db()
        self.assertEqual(later.amount_paid, Decimal('500'))
        self.assertEqual(later.balance, Decimal('300'))
        self.assertEqual(later.payment_status, 'partial')
        self.assertEqual(credit_balance('Kibo Traders'), Decimal('0'))

        # The draw is recorded as a credit-funded payment, not new money
        drawn = later.payments.get()
        self.assertTrue(drawn.from_credit)

    def test_credit_use_is_capped_by_what_is_owed(self):
        self.receive(self.sale, '3000')  # 2,000 on account
        small = CustomerRecord.objects.create(
            date=date(2026, 9, 1), customer_name='Kibo Traders',
            receipt_number='RC-3', sales_amount=300)

        self.client.post(f'/api/crm-records/{small.id}/apply_credit/', {}, content_type='application/json')
        small.refresh_from_db()
        self.assertEqual(small.balance, Decimal('0'))
        self.assertEqual(credit_balance('Kibo Traders'), Decimal('1700'))

    def test_cannot_spend_more_credit_than_is_held(self):
        deposit = self.client.post('/api/crm-credits/', {
            'customer_name': 'Kibo Traders', 'amount': '200',
            'received_on': '2026-08-01', 'method': 'cash',
        }, content_type='application/json')
        self.assertEqual(deposit.status_code, 201, deposit.content)

        response = self.client.post(f'/api/crm-records/{self.sale.id}/apply_credit/',
                                    {'amount': '500'}, content_type='application/json')
        self.assertEqual(response.status_code, 400)
        self.assertIn('Only 200.00 of credit is available', response.json()['detail'])
        self.assertEqual(credit_balance('Kibo Traders'), Decimal('200'))

    def test_apply_credit_needs_credit_and_an_unpaid_sale(self):
        no_credit = self.client.post(f'/api/crm-records/{self.sale.id}/apply_credit/',
                                     {}, content_type='application/json')
        self.assertEqual(no_credit.status_code, 400)
        self.assertIn('no credit on account', no_credit.json()['detail'])

        self.receive(self.sale, '1000')
        self.client.post('/api/crm-credits/', {
            'customer_name': 'Kibo Traders', 'amount': '100', 'received_on': '2026-08-01',
        }, content_type='application/json')
        settled = self.client.post(f'/api/crm-records/{self.sale.id}/apply_credit/',
                                   {}, content_type='application/json')
        self.assertEqual(settled.status_code, 400)
        self.assertIn('already fully paid', settled.json()['detail'])

    def test_deposit_before_any_order(self):
        """Customer pays in advance; the money waits on their account."""
        response = self.client.post('/api/crm-credits/', {
            'customer_name': 'Walk In Ltd', 'amount': '750',
            'received_on': '2026-08-01', 'method': 'mobile', 'reference': 'MPESA-9',
        }, content_type='application/json')
        self.assertEqual(response.status_code, 201, response.content)
        self.assertEqual(credit_balance('Walk In Ltd'), Decimal('750'))

        balance = self.client.get('/api/crm-credits/balance/', {'customer': 'Walk In Ltd'}).json()
        self.assertEqual(Decimal(str(balance['credit_available'])), Decimal('750'))
        self.assertEqual(len(balance['entries']), 1)

    def test_spent_credit_cannot_be_deleted(self):
        self.receive(self.sale, '1500')
        later = CustomerRecord.objects.create(date=date(2026, 9, 1), customer_name='Kibo Traders',
                                              receipt_number='RC-9', sales_amount=800)
        self.client.post(f'/api/crm-records/{later.id}/apply_credit/', {}, content_type='application/json')

        credit = CrmCredit.objects.get()
        response = self.client.delete(f'/api/crm-credits/{credit.id}/')
        self.assertEqual(response.status_code, 400)
        self.assertIn('cannot be removed', response.json()['detail'])
        self.assertEqual(CrmCredit.objects.count(), 1)

    def test_unspent_credit_can_be_deleted(self):
        self.receive(self.sale, '1500')
        credit = CrmCredit.objects.get()
        self.assertEqual(self.client.delete(f'/api/crm-credits/{credit.id}/').status_code, 204)
        self.assertEqual(credit_balance('Kibo Traders'), Decimal('0'))

    def test_credit_shows_on_the_customer_list_and_summary(self):
        self.receive(self.sale, '1500')
        row = self.client.get('/api/crm-records/customers/').json()['results'][0]
        self.assertEqual(Decimal(str(row['credit_available'])), Decimal('500'))

        summary = self.client.get('/api/crm-records/summary/').json()
        self.assertEqual(Decimal(str(summary['credit_available'])), Decimal('500'))

    def test_receive_rejects_junk(self):
        self.assertEqual(self.receive(self.sale, '0').status_code, 400)
        self.assertEqual(self.receive(self.sale, '-100').status_code, 400)
        self.assertEqual(self.receive(self.sale, 'abc').status_code, 400)

    def test_credit_endpoints_are_closed_to_other_roles(self):
        self.client.force_login(User.objects.create_user(username='crm_rep6', password='pw', role='sales_rep'))
        self.assertEqual(self.receive(self.sale, '100').status_code, 403)
        self.assertEqual(
            self.client.post(f'/api/crm-records/{self.sale.id}/apply_credit/').status_code, 403)
        self.assertEqual(self.client.get('/api/crm-credits/').status_code, 403)


class CrmPaymentAggregationTest(TestCase):
    """Payment sums must not corrupt the sales totals they sit beside."""

    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct8', password='pw', role='accountant')
        self.client.force_login(self.accountant)
        self.a = CustomerRecord.objects.create(date=date(2026, 8, 1), customer_name='Kibo Traders',
                                               receipt_number='RC-1', sales_amount=1000)
        self.b = CustomerRecord.objects.create(date=date(2026, 8, 2), customer_name='Kibo Traders',
                                               receipt_number='RC-2', sales_amount=500)
        # Three payments on one sale is where a naive join would triple its
        # sales_amount in the per-customer totals.
        for amount in (100, 200, 300):
            CrmPayment.objects.create(record=self.a, amount=amount, paid_on=date(2026, 8, 3))

    def test_customer_totals_are_not_inflated_by_multiple_payments(self):
        row = self.client.get('/api/crm-records/customers/').json()['results'][0]
        self.assertEqual(row['records'], 2)
        self.assertEqual(float(row['total_amount']), 1500.0)   # not 1000*3 + 500
        self.assertEqual(float(row['amount_paid']), 600.0)
        self.assertEqual(float(row['balance']), 900.0)
        self.assertEqual(row['payment_status'], 'partial')

    def test_summary_is_not_inflated_either(self):
        summary = self.client.get('/api/crm-records/summary/').json()
        self.assertEqual(summary['records'], 2)
        self.assertEqual(float(summary['total_amount']), 1500.0)
        self.assertEqual(float(summary['amount_paid']), 600.0)

    def test_status_filter(self):
        CustomerRecord.objects.create(date=date(2026, 8, 4), customer_name='Paid Customer',
                                      receipt_number='RC-3', sales_amount=200)
        paid_record = CustomerRecord.objects.get(receipt_number='RC-3')
        CrmPayment.objects.create(record=paid_record, amount=200, paid_on=date(2026, 8, 4))

        def receipts(status):
            body = self.client.get('/api/crm-records/', {'status': status}).json()
            return sorted(r['receipt_number'] for r in body['results'])

        self.assertEqual(receipts('unpaid'), ['RC-2'])
        self.assertEqual(receipts('partial'), ['RC-1'])
        self.assertEqual(receipts('paid'), ['RC-3'])

    def test_csv_export_carries_payment_columns(self):
        body = self.client.get(reverse('crm:export')).content.decode()
        header = body.splitlines()[0]
        self.assertIn('Amount Paid', header)
        self.assertIn('Balance', header)
        self.assertIn('Status', header)
        self.assertIn('partial', body)


class CrmPdfReportTest(TestCase):
    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct6', password='pw', role='accountant')
        self.client.force_login(self.accountant)
        CustomerRecord.objects.create(date=date(2026, 8, 1), customer_name='Kibo Traders',
                                      receipt_number='RC-1', tin='109-882-441', sales_amount=100)
        CustomerRecord.objects.create(date=date(2026, 8, 5), customer_name='Kibo Traders',
                                      receipt_number='RC-2', tin='109-882-441', sales_amount=250)
        CustomerRecord.objects.create(date=date(2026, 8, 3), customer_name='Mwanza Const',
                                      receipt_number='RC-3', tin='122-004-908', sales_amount=900)

    def assertIsPdf(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF-'), 'response is not a PDF')
        self.assertGreater(len(response.content), 1000)

    def test_all_customers_report(self):
        response = self.client.get(reverse('crm:report'))
        self.assertIsPdf(response)
        self.assertIn('crm_customers_', response['Content-Disposition'])

    def test_all_customers_report_honours_filters(self):
        filtered = self.client.get(reverse('crm:report'), {'search': 'Mwanza'})
        self.assertIsPdf(filtered)
        # A filter that matches nothing still renders a valid (empty) report
        empty = self.client.get(reverse('crm:report'), {'search': 'nobody at all'})
        self.assertIsPdf(empty)

    def test_single_customer_report(self):
        response = self.client.get(reverse('crm:customer_report'), {'name': 'Kibo Traders'})
        self.assertIsPdf(response)
        self.assertIn('crm_Kibo_Traders_', response['Content-Disposition'])

    def test_single_customer_report_needs_a_known_customer(self):
        self.assertEqual(self.client.get(reverse('crm:customer_report')).status_code, 404)
        self.assertEqual(
            self.client.get(reverse('crm:customer_report'), {'name': 'Nobody Ltd'}).status_code, 404)

    def test_reports_carry_the_logo(self):
        """The logo must actually be embedded, not silently dropped."""
        from apps.crm.reports import _logo_path
        self.assertIsNotNone(_logo_path(), 'no logo file resolved')

        response = self.client.get(reverse('crm:customer_report'), {'name': 'Kibo Traders'})
        self.assertIn(b'/Image', response.content, 'no image XObject in the PDF')

    def test_reports_are_closed_to_other_roles(self):
        self.client.force_login(User.objects.create_user(username='crm_rep4', password='pw', role='sales_rep'))
        self.assertEqual(self.client.get(reverse('crm:report')).status_code, 403)
        self.assertEqual(
            self.client.get(reverse('crm:customer_report'), {'name': 'Kibo Traders'}).status_code, 403)


class CrmPaginationTest(TestCase):
    """Neither table may ever return the whole register."""

    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct5', password='pw', role='accountant')
        self.client.force_login(self.accountant)
        CustomerRecord.objects.bulk_create([
            CustomerRecord(date=date(2026, 8, 1), customer_name=f'Customer {i:03d}',
                           receipt_number=f'RC-{i:03d}', sales_amount=10 * i)
            for i in range(1, 46)
        ])

    def test_records_default_to_20_per_page(self):
        body = self.client.get('/api/crm-records/').json()
        self.assertEqual(body['count'], 45)
        self.assertEqual(len(body['results']), 20)
        self.assertIsNotNone(body['next'])

    def test_page_size_options(self):
        for size in (10, 20, 50, 100):
            body = self.client.get('/api/crm-records/', {'page_size': size}).json()
            self.assertEqual(len(body['results']), min(size, 45), f'page_size={size}')

    def test_page_size_is_capped_at_100(self):
        body = self.client.get('/api/crm-records/', {'page_size': 5000}).json()
        self.assertEqual(len(body['results']), 45)  # capped, not unbounded
        CustomerRecord.objects.bulk_create([
            CustomerRecord(date=date(2026, 8, 2), customer_name=f'Extra {i}', sales_amount=1)
            for i in range(120)
        ])
        body = self.client.get('/api/crm-records/', {'page_size': 5000}).json()
        self.assertEqual(len(body['results']), 100)

    def test_paging_walks_the_whole_set_without_repeats(self):
        seen, page = [], 1
        while True:
            body = self.client.get('/api/crm-records/', {'page_size': 10, 'page': page}).json()
            seen.extend(r['id'] for r in body['results'])
            if not body['next']:
                break
            page += 1
        self.assertEqual(len(seen), 45)
        self.assertEqual(len(set(seen)), 45)

    def test_customer_list_is_paginated_too(self):
        body = self.client.get('/api/crm-records/customers/', {'page_size': 10}).json()
        self.assertEqual(body['count'], 45)
        self.assertEqual(len(body['results']), 10)

    def test_summary_totals_the_whole_selection_not_the_page(self):
        """The stat tiles must not report only what is on screen."""
        summary = self.client.get('/api/crm-records/summary/', {'page_size': 10}).json()
        self.assertEqual(summary['records'], 45)
        self.assertEqual(summary['customers'], 45)
        self.assertEqual(float(summary['total_amount']), sum(10 * i for i in range(1, 46)))

    def test_csv_export_is_not_paginated(self):
        """Export gives every matching row, not the current page."""
        response = self.client.get(reverse('crm:export'), {'page_size': 10})
        rows = response.content.decode().strip().splitlines()
        self.assertEqual(len(rows), 46)  # 45 records + header


def build_xlsx(rows, headers=None, name='customers.xlsx'):
    """An in-memory .xlsx upload, as the browser would send it."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers if headers is not None
                 else ['Date', 'Receipt No', 'EFD Receipt No', 'Customer Name', 'TIN',
                       'Sales Amount', 'Amount Paid'])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        name, buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


class CrmBulkUploadTest(TestCase):
    def setUp(self):
        self.accountant = User.objects.create_user(username='crm_acct3', password='pw', role='accountant')
        self.client.force_login(self.accountant)

    def post_file(self, upload):
        return self.client.post('/api/crm-records/bulk_upload/', {'file': upload})

    def test_uploads_rows_from_excel(self):
        upload = build_xlsx([
            ['2026-08-01', 'RC-0012', '35EFD9921', 'Kibo Traders', '109-882-441', 1250000],
            [date(2026, 8, 3), 'RC-0013', '35EFD9944', 'Mwanza Const', '122-004-908', '840,500.50'],
        ])
        response = self.post_file(upload)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['created'], 2)

        first = CustomerRecord.objects.get(receipt_number='RC-0012')
        self.assertEqual(first.date, date(2026, 8, 1))
        self.assertEqual(first.customer_name, 'Kibo Traders')
        self.assertEqual(first.tin, '109-882-441')
        self.assertEqual(first.sales_amount, Decimal('1250000.00'))
        self.assertEqual(first.created_by, self.accountant)

        # A real date cell and a comma-formatted amount both survive
        second = CustomerRecord.objects.get(receipt_number='RC-0013')
        self.assertEqual(second.date, date(2026, 8, 3))
        self.assertEqual(second.sales_amount, Decimal('840500.50'))

    def test_headers_are_matched_loosely_and_extras_ignored(self):
        upload = build_xlsx(
            [['01/08/2026', 'Kibo Traders', '35EFD9921', '109-882-441', 500, 'ignored']],
            headers=['  DATE ', 'customer', 'EFD Receipt Number', 'TIN No', 'Amount', 'Notes'],
        )
        response = self.post_file(upload)
        self.assertEqual(response.status_code, 200, response.content)

        record = CustomerRecord.objects.get()
        self.assertEqual(record.date, date(2026, 8, 1))
        self.assertEqual(record.customer_name, 'Kibo Traders')
        self.assertEqual(record.efd_receipt_number, '35EFD9921')
        self.assertEqual(record.sales_amount, Decimal('500'))

    def test_bad_rows_are_reported_and_good_rows_still_import(self):
        upload = build_xlsx([
            ['2026-08-01', 'RC-1', '', 'Good Customer', '', 100],
            ['not a date', 'RC-2', '', 'Bad Date', '', 100],
            ['2026-08-02', 'RC-3', '', '', '', 100],
            [],  # blank spacer row is skipped silently
            ['2026-08-03', 'RC-4', '', 'Bad Amount', '', 'abc'],
            ['2026-08-04', 'RC-5', '', 'Blank Amount', '', ''],
        ])
        response = self.post_file(upload)
        self.assertEqual(response.status_code, 207)
        body = response.json()
        self.assertEqual(body['created'], 2)  # Good Customer + Blank Amount
        self.assertEqual(body['error_count'], 3)
        for row_label in ('Row 3:', 'Row 4:', 'Row 6:'):
            self.assertTrue(any(e.startswith(row_label) for e in body['errors']), body['errors'])
        self.assertEqual(CustomerRecord.objects.filter(customer_name__in=['Bad Date', 'Bad Amount']).count(), 0)
        # A blank amount cell is not an error — it just means nothing recorded yet
        self.assertEqual(CustomerRecord.objects.get(receipt_number='RC-5').sales_amount, Decimal('0.00'))

    def test_reupload_skips_rows_already_in_register(self):
        rows = [['2026-08-01', 'RC-0012', '35EFD9921', 'Kibo Traders', '', 100]]
        self.assertEqual(self.post_file(build_xlsx(rows)).json()['created'], 1)

        again = self.post_file(build_xlsx(rows)).json()
        self.assertEqual(again['created'], 0)
        self.assertEqual(again['skipped'], 1)
        self.assertEqual(CustomerRecord.objects.count(), 1)

    def test_missing_required_column_is_rejected(self):
        upload = build_xlsx([['RC-1', 100]], headers=['Receipt No', 'Amount'])
        response = self.post_file(upload)
        self.assertEqual(response.status_code, 400)
        self.assertIn('Customer Name', response.json()['detail'])
        self.assertEqual(CustomerRecord.objects.count(), 0)

    def test_wrong_file_type_is_rejected(self):
        upload = SimpleUploadedFile('customers.pdf', b'%PDF-1.4 not a spreadsheet', content_type='application/pdf')
        response = self.post_file(upload)
        self.assertEqual(response.status_code, 400)
        self.assertIn('Unsupported file type', response.json()['detail'])

    def test_csv_is_accepted_too(self):
        csv_bytes = (
            'Date,Receipt No,EFD Receipt No,Customer Name,TIN,Sales Amount\n'
            '2026-08-01,RC-0012,35EFD9921,Kibo Traders,109-882-441,1250000\n'
        ).encode('utf-8')
        response = self.post_file(SimpleUploadedFile('customers.csv', csv_bytes, content_type='text/csv'))
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(CustomerRecord.objects.get().customer_name, 'Kibo Traders')

    def test_amount_paid_column_creates_opening_payments(self):
        upload = build_xlsx([
            ['2026-08-01', 'RC-1', '', 'Kibo Traders', '', 1000, 1000],   # settled
            ['2026-08-02', 'RC-2', '', 'Kibo Traders', '', 1000, 400],    # part paid
            ['2026-08-03', 'RC-3', '', 'Kibo Traders', '', 1000, ''],     # nothing paid
        ])
        response = self.post_file(upload)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()['payments'], 2)

        by_receipt = {r.receipt_number: r for r in CustomerRecord.objects.all()}
        self.assertEqual(by_receipt['RC-1'].payment_status, 'paid')
        self.assertEqual(by_receipt['RC-2'].amount_paid, Decimal('400'))
        self.assertEqual(by_receipt['RC-2'].payment_status, 'partial')
        self.assertEqual(by_receipt['RC-3'].payment_status, 'unpaid')

        opening = by_receipt['RC-2'].payments.get()
        self.assertEqual(opening.paid_on, date(2026, 8, 2))  # dated with the sale
        self.assertEqual(opening.reference, 'Imported opening balance')

    def test_imported_overpayment_goes_to_customer_credit(self):
        upload = build_xlsx([['2026-08-01', 'RC-1', '', 'Kibo Traders', '', 1000, 1500]])
        response = self.post_file(upload)
        body = response.json()
        self.assertEqual(body['payments'], 1)
        self.assertEqual(body['credits'], 1)

        record = CustomerRecord.objects.get()
        self.assertEqual(record.amount_paid, Decimal('1000'))  # capped at the sale
        self.assertEqual(record.payment_status, 'paid')
        self.assertEqual(credit_balance('Kibo Traders'), Decimal('500'))

    def test_amount_paid_header_variants_and_absence(self):
        upload = build_xlsx(
            [['2026-08-01', 'Kibo Traders', 1000, 250]],
            headers=['Date', 'Customer Name', 'Sales Amount', 'PAID'],
        )
        self.assertEqual(self.post_file(upload).status_code, 200)
        self.assertEqual(CustomerRecord.objects.get().amount_paid, Decimal('250'))

        # A sheet with no paid column still imports, as unpaid
        CustomerRecord.objects.all().delete()
        upload = build_xlsx(
            [['2026-08-01', 'Mwanza Const', 500]],
            headers=['Date', 'Customer Name', 'Sales Amount'],
        )
        self.assertEqual(self.post_file(upload).status_code, 200)
        self.assertEqual(CustomerRecord.objects.get().payment_status, 'unpaid')

    def test_negative_amount_paid_is_rejected_per_row(self):
        upload = build_xlsx([
            ['2026-08-01', 'RC-1', '', 'Good', '', 1000, 100],
            ['2026-08-02', 'RC-2', '', 'Bad', '', 1000, -50],
        ])
        response = self.post_file(upload)
        self.assertEqual(response.status_code, 207)
        self.assertEqual(response.json()['created'], 1)
        self.assertTrue(any('Amount Paid cannot be negative' in e for e in response.json()['errors']))

    def test_no_file_returns_400(self):
        response = self.client.post('/api/crm-records/bulk_upload/', {})
        self.assertEqual(response.status_code, 400)

    def test_template_download_is_a_readable_workbook(self):
        response = self.client.get(reverse('crm:import_template'))
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        sheet = load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(
            [c.value for c in sheet[1]],
            ['Date', 'Receipt No', 'EFD Receipt No', 'Customer Name', 'TIN',
             'Sales Amount', 'Amount Paid'],
        )

    def test_upload_is_closed_to_other_roles(self):
        self.client.force_login(User.objects.create_user(username='crm_rep2', password='pw', role='sales_rep'))
        response = self.post_file(build_xlsx([['2026-08-01', 'RC-9', '', 'Sneaky', '', 1]]))
        self.assertEqual(response.status_code, 403)
        self.assertEqual(CustomerRecord.objects.count(), 0)
        self.assertEqual(self.client.get(reverse('crm:import_template')).status_code, 403)


class CrmSalesSyncTest(TestCase):
    """A sale made at the POS lands in the CRM register carrying its payment
    state — paid in full, part paid, or wholly on credit — and keeps up as
    money comes in against it."""

    def setUp(self):
        self.branch = Branch.objects.create(name='Sync Branch')
        self.cashier = User.objects.create_user(username='sync_rep', password='pw', role='sales_rep')
        self.customer = Customer.objects.create(name='Mbeya Traders', phone='0700111222')

    def make_sale(self, total, invoice, customer=None, status='pending'):
        return Sale.objects.create(
            invoice_number=invoice, branch=self.branch, user=self.cashier,
            customer=customer, customer_name='' if customer else 'Walk-in Customer',
            total_amount=Decimal(total), status=status,
        )

    def record_for(self, sale):
        return CustomerRecord.objects.filter(source_invoice=sale.invoice_number).first()

    def pay(self, sale, amount, method='cash', reference=''):
        return Transaction.objects.create(
            sale=sale, amount=Decimal(amount), payment_method=method, reference=reference)

    # --- the three states the register must show -------------------------

    def test_credit_sale_appears_unpaid(self):
        sale = self.make_sale('500000', 'INV-CREDIT', self.customer)
        record = self.record_for(sale)
        self.assertIsNotNone(record, 'the sale should have created a CRM row')
        self.assertEqual(record.customer_name, 'Mbeya Traders')
        self.assertEqual(record.sales_amount, Decimal('500000'))
        self.assertEqual(record.amount_paid, Decimal('0'))
        self.assertEqual(record.payment_status, 'unpaid')

    def test_deposit_makes_it_partial(self):
        sale = self.make_sale('500000', 'INV-PART', self.customer)
        self.pay(sale, '200000', 'mobile')
        record = self.record_for(sale)
        self.assertEqual(record.amount_paid, Decimal('200000'))
        self.assertEqual(record.balance, Decimal('300000'))
        self.assertEqual(record.payment_status, 'partial')

    def test_full_payment_makes_it_paid(self):
        sale = self.make_sale('120000', 'INV-PAID')
        self.pay(sale, '120000')
        record = self.record_for(sale)
        self.assertEqual(record.payment_status, 'paid')
        self.assertEqual(record.customer_name, 'Walk-in Customer')

    def test_later_payment_moves_it_from_credit_to_paid(self):
        """Collections recorded on the Debtors screen reach the register."""
        sale = self.make_sale('80000', 'INV-COLLECT', self.customer)
        self.assertEqual(self.record_for(sale).payment_status, 'unpaid')
        self.pay(sale, '30000', 'bank', reference='SLIP-9')
        self.assertEqual(self.record_for(sale).payment_status, 'partial')
        self.pay(sale, '50000', 'cash')
        record = self.record_for(sale)
        self.assertEqual(record.payment_status, 'paid')
        self.assertEqual(record.balance, Decimal('0'))

    def test_payment_method_and_reference_carry_across(self):
        sale = self.make_sale('10000', 'INV-METHOD', self.customer)
        self.pay(sale, '10000', 'bank', reference='SLIP-42')
        payment = self.record_for(sale).payments.get()
        self.assertEqual(payment.method, 'bank')
        self.assertEqual(payment.reference, 'SLIP-42')

    # --- it must not double count ----------------------------------------

    def test_resaving_the_sale_does_not_duplicate_anything(self):
        sale = self.make_sale('90000', 'INV-RESAVE', self.customer)
        self.pay(sale, '40000')
        for _ in range(3):
            sale.save()
        self.assertEqual(CustomerRecord.objects.filter(source_invoice='INV-RESAVE').count(), 1)
        record = self.record_for(sale)
        self.assertEqual(record.payments.count(), 1)
        self.assertEqual(record.amount_paid, Decimal('40000'))

    def test_backfill_command_is_repeatable(self):
        sale = self.make_sale('75000', 'INV-BACKFILL', self.customer)
        self.pay(sale, '25000')
        for _ in range(2):
            call_command('sync_crm_from_sales', verbosity=0)
        self.assertEqual(CustomerRecord.objects.filter(source_invoice='INV-BACKFILL').count(), 1)
        self.assertEqual(self.record_for(sale).amount_paid, Decimal('25000'))

    def test_amount_follows_the_sale_total(self):
        sale = self.make_sale('100000', 'INV-AMEND', self.customer)
        sale.total_amount = Decimal('85000')
        sale.save()
        self.assertEqual(self.record_for(sale).sales_amount, Decimal('85000'))

    # --- hand-kept work is never trampled --------------------------------

    def test_hand_entered_details_survive_a_resync(self):
        sale = self.make_sale('60000', 'INV-HAND', self.customer)
        record = self.record_for(sale)
        record.tin = '109-882-441'
        record.efd_receipt_number = '35EFD9921'
        record.save()

        sale.save()
        call_command('sync_crm_from_sales', verbosity=0)

        record.refresh_from_db()
        self.assertEqual(record.tin, '109-882-441')
        self.assertEqual(record.efd_receipt_number, '35EFD9921')

    def test_a_hand_typed_payment_is_left_alone(self):
        sale = self.make_sale('70000', 'INV-MANUAL', self.customer)
        record = self.record_for(sale)
        CrmPayment.objects.create(record=record, amount=Decimal('10000'), paid_on=date(2026, 9, 1))
        self.pay(sale, '20000')
        call_command('sync_crm_from_sales', verbosity=0)

        record.refresh_from_db()
        self.assertEqual(record.payments.count(), 2)
        self.assertEqual(record.amount_paid, Decimal('30000'))

    def test_reversing_a_payment_reverses_it_here(self):
        sale = self.make_sale('50000', 'INV-REVERSE', self.customer)
        txn = self.pay(sale, '50000')
        self.assertEqual(self.record_for(sale).payment_status, 'paid')
        txn.delete()
        self.assertEqual(self.record_for(sale).payment_status, 'unpaid')

    # --- cancelled sales --------------------------------------------------

    def test_cancelled_sale_drops_out_of_the_register(self):
        sale = self.make_sale('40000', 'INV-CANCEL', self.customer)
        self.assertIsNotNone(self.record_for(sale))
        sale.status = 'cancelled'
        sale.save()
        self.assertIsNone(self.record_for(sale))

    def test_cancelled_sale_is_kept_once_someone_has_worked_on_it(self):
        sale = self.make_sale('40000', 'INV-CANCEL2', self.customer)
        record = self.record_for(sale)
        record.efd_receipt_number = '35EFD0001'
        record.save()
        sale.status = 'cancelled'
        sale.save()
        self.assertIsNotNone(self.record_for(sale), 'a row with EFD details must not vanish')

    def test_manual_records_are_untouched_by_the_sync(self):
        manual = CustomerRecord.objects.create(
            date=date(2026, 8, 1), customer_name='Walk-in cash buyer',
            sales_amount=Decimal('15000'))
        self.make_sale('40000', 'INV-OTHER', self.customer)
        call_command('sync_crm_from_sales', verbosity=0)
        manual.refresh_from_db()
        self.assertEqual(manual.sales_amount, Decimal('15000'))
        self.assertEqual(manual.source_invoice, '')
