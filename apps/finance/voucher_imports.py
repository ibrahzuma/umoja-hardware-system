"""Bulk upload of vouchers (and ledgers) into the books from Excel or CSV.

Back-dated work arrives as a spreadsheet, not one voucher at a time. The
workbook the accountant fills in has:

* a **Vouchers** sheet — one row per *line*, with a **Ref** that groups the
  lines of one voucher (any text: "1", "2", the invoice number…). Debit and
  Credit rows of the same Ref, Type and Date become one voucher, so a
  three-line VAT sale is three rows sharing a Ref;
* an optional **Ledgers** sheet — ledgers to create first (or opening
  balances to set on ones that exist), so a file can carry a whole set of
  books from day one.

Nothing here posts by a different route: every voucher goes through
`vouchers.post_voucher` with exactly the checks the form gets, and the whole
file is posted in one transaction — either every voucher lands or none does,
and the report says which rows to fix. `run(commit=False)` is the same pass
with the posting rolled back, which is what the "Check file" button shows.
"""

import csv
import io
import re
from collections import OrderedDict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError

from apps.inventory.models import Supplier
from apps.sales.models import Customer
from .models import LedgerAccount, Voucher
from . import vouchers

MAX_ROWS = 5000
EXCEL_EPOCH = datetime(1899, 12, 30)
DATE_FORMATS = ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%Y/%m/%d', '%d %b %Y', '%d %B %Y')


class ImportError_(Exception):
    """The file itself cannot be used (wrong type, no headers, missing columns)."""


# ---------------------------------------------------------------------------
# Columns
# ---------------------------------------------------------------------------

VOUCHER_HEADERS = ['Ref', 'Type', 'Date', 'Ledger', 'Side', 'Amount', 'Description',
                   'Party', 'Invoice No', 'EFD No', 'Against', 'Narration']
LEDGER_HEADERS = ['Code', 'Name', 'Kind', 'Opening Balance', 'Opening Side', 'Notes']

VOUCHER_ALIASES = {
    'ref': 'ref', 'voucherref': 'ref', 'voucher': 'ref', 'voucherno': 'ref', 'no': 'ref', 'group': 'ref',
    'type': 'type', 'vouchertype': 'type',
    'date': 'date', 'transactiondate': 'date', 'dateoftransaction': 'date', 'tarehe': 'date',
    'ledger': 'ledger', 'account': 'ledger', 'ledgeraccount': 'ledger', 'accountname': 'ledger',
    'ledgercode': 'ledger', 'code': 'ledger',
    'side': 'side', 'drcr': 'side', 'debitcredit': 'side',
    'amount': 'amount', 'kiasi': 'amount',
    'description': 'description', 'maelezo': 'description', 'details': 'description',
    'party': 'party', 'customer': 'party', 'supplier': 'party', 'customersupplier': 'party',
    'mteja': 'party',
    'invoiceno': 'invoice_number', 'invoice': 'invoice_number', 'invoicenumber': 'invoice_number',
    'salesinvoiceno': 'invoice_number', 'purchaseinvoiceno': 'invoice_number',
    'efdno': 'efd_number', 'efd': 'efd_number', 'efdrctno': 'efd_number', 'efdreceiptno': 'efd_number',
    'rctno': 'efd_number',
    'against': 'against', 'allocateto': 'against', 'againstinvoice': 'against', 'settles': 'against',
    'narration': 'narration', 'linenote': 'narration', 'note': 'narration',
}
# What a row may write in the Ledger column to mean "the party's own ledger".
PARTY_WORDS = ('customer', 'debtor', 'supplier', 'creditor', 'party')
LEDGER_ALIASES = {
    'code': 'code', 'ledgercode': 'code',
    'name': 'name', 'ledger': 'name', 'ledgername': 'name', 'account': 'name',
    'kind': 'kind', 'type': 'kind', 'category': 'kind',
    'openingbalance': 'opening_balance', 'opening': 'opening_balance', 'balance': 'opening_balance',
    'openingside': 'opening_side', 'side': 'opening_side', 'drcr': 'opening_side',
    'notes': 'notes', 'note': 'notes',
}

TYPE_ALIASES = {
    'sales': 'sales', 'sale': 'sales', 'sv': 'sales', 'salesvoucher': 'sales', 'mauzo': 'sales',
    'purchase': 'purchase', 'purchases': 'purchase', 'pu': 'purchase', 'purchasevoucher': 'purchase',
    'manunuzi': 'purchase',
    'receipt': 'receipt', 'receipts': 'receipt', 'rv': 'receipt', 'rcpt': 'receipt',
    'payment': 'payment', 'payments': 'payment', 'pv': 'payment', 'pay': 'payment', 'malipo': 'payment',
    'contra': 'contra', 'cv': 'contra', 'transfer': 'contra',
    'journal': 'journal', 'jv': 'journal', 'jnl': 'journal',
}
SIDE_ALIASES = {
    'dr': 'debit', 'debit': 'debit', 'd': 'debit',
    'cr': 'credit', 'credit': 'credit', 'c': 'credit',
}
KIND_ALIASES = {k: k for k, _ in LedgerAccount.KINDS}
KIND_ALIASES.update({
    'bankaccount': 'bank', 'cashbook': 'cash', 'pettycash': 'cash',
    'debtor': 'customer', 'receivable': 'customer', 'creditor': 'supplier', 'payable': 'supplier',
    'revenue': 'income', 'sales': 'income', 'expenses': 'expense', 'otherasset': 'asset',
    'otherliability': 'liability', 'vat': 'tax', 'capital': 'equity',
})


def _normalise(text):
    return re.sub(r'[^a-z0-9]', '', str(text or '').lower())


def _map_headers(raw_headers, aliases, required, expected):
    mapping = {}
    for index, header in enumerate(raw_headers):
        field = aliases.get(_normalise(header))
        if field and field not in mapping.values():
            mapping[index] = field
    missing = [f for f in required if f not in mapping.values()]
    if missing:
        raise ImportError_(
            'The sheet is missing required column(s): ' + ', '.join(missing)
            + '. Expected headers: ' + ', '.join(expected) + '.')
    return mapping


def parse_date(value):
    if value in (None, ''):
        raise ValueError('Date is required')
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return (EXCEL_EPOCH + timedelta(days=float(value))).date()
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f'"{text}" is not a date (use YYYY-MM-DD or DD/MM/YYYY)')


def parse_amount(value):
    if value in (None, ''):
        return Decimal('0.00')
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal('0.01'))
    text = re.sub(r'[^0-9.\-]', '', str(value).strip())
    if text in ('', '-', '.'):
        raise ValueError(f'"{value}" is not a number')
    try:
        return Decimal(text).quantize(Decimal('0.01'))
    except InvalidOperation:
        raise ValueError(f'"{value}" is not a number')


def _text(value, limit=200):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()[:limit]


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------

def _read_sheets(upload):
    """{sheet name (normalised): [rows]} for an .xlsx, or {'vouchers': rows}
    for a .csv, which can only ever carry the one sheet."""
    name = (getattr(upload, 'name', '') or '').lower()
    data = upload.read()
    if not data:
        raise ImportError_('The uploaded file is empty.')

    if name.endswith(('.csv', '.txt')):
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('latin-1')
        return {'vouchers': list(csv.reader(io.StringIO(text)))}

    if not name.endswith(('.xlsx', '.xlsm')):
        raise ImportError_('Upload an Excel .xlsx file or a .csv. Old .xls files must be re-saved as .xlsx.')
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportError_('That file could not be read as an Excel workbook.')
    try:
        sheets = {}
        for ws in workbook.worksheets:
            sheets[_normalise(ws.title)] = list(ws.iter_rows(values_only=True))
        return sheets
    finally:
        workbook.close()


def _header_and_rows(rows):
    """Skip leading blank rows; return (header, [(excel_row_number, row)])."""
    for i, candidate in enumerate(rows):
        if candidate and any(_normalise(c) for c in candidate):
            body = [(i + 2 + n, r) for n, r in enumerate(rows[i + 1:])
                    if r and any(str(c).strip() for c in r if c is not None)]
            return candidate, body
    return None, []


def _pick_voucher_sheet(sheets):
    """The Vouchers sheet by name, else the first sheet that is not Ledgers."""
    for key in ('vouchers', 'voucher', 'entries', 'transactions'):
        if key in sheets:
            return sheets[key]
    for key, rows in sheets.items():
        if key not in ('ledgers', 'ledgerlist', 'chartofaccounts', 'instructions', 'readme'):
            return rows
    return []


# ---------------------------------------------------------------------------
# Ledgers sheet
# ---------------------------------------------------------------------------

def _parse_ledgers(rows):
    header, body = _header_and_rows(rows)
    if header is None:
        return []
    mapping = _map_headers(header, LEDGER_ALIASES, ['name', 'kind'], LEDGER_HEADERS)
    out = []
    for row_no, raw in body[:MAX_ROWS]:
        rec = {'row': row_no, 'code': '', 'name': '', 'kind': '', 'opening_balance': None,
               'opening_side': '', 'notes': ''}
        for index, field in mapping.items():
            if index < len(raw):
                rec[field] = raw[index]
        rec['code'] = _text(rec['code'], 20).upper()
        rec['name'] = _text(rec['name'], 150)
        rec['notes'] = _text(rec['notes'], 500)
        out.append(rec)
    return out


def _apply_ledgers(records, report):
    """Create the ledgers the sheet names (or set opening balances on ones
    that exist). Errors go on the report; nothing raises."""
    for rec in records:
        where = f"Ledgers row {rec['row']}"
        try:
            kind = KIND_ALIASES.get(_normalise(rec['kind']))
            if not kind:
                raise ValueError(f"kind \"{rec['kind']}\" is not one of: "
                                 + ', '.join(k for k, _ in LedgerAccount.KINDS))
            if not rec['name']:
                raise ValueError('the ledger needs a name')
            if kind in ('customer', 'supplier'):
                raise ValueError('customer and supplier ledgers are made from the customer or supplier '
                                 'itself — name the party on a voucher row instead')
            side_raw = rec['opening_side']
            side = SIDE_ALIASES.get(_normalise(side_raw)) if _text(side_raw) else ''
            if _text(side_raw) and not side:
                raise ValueError(f'opening side "{side_raw}" must be Dr or Cr')
            opening = None
            if rec['opening_balance'] not in (None, ''):
                opening = parse_amount(rec['opening_balance'])
                if opening < 0:
                    raise ValueError('use the side (Dr/Cr), not a minus sign, for the opening balance')

            ledger = None
            if rec['code']:
                ledger = LedgerAccount.objects.filter(code__iexact=rec['code']).first()
            if ledger is None:
                ledger = LedgerAccount.objects.filter(kind=kind, name__iexact=rec['name']).first()

            if ledger is None:
                ledger = LedgerAccount(kind=kind, name=rec['name'], code=rec['code'], notes=rec['notes'])
                if opening is not None:
                    ledger.opening_balance = opening
                    ledger.opening_side = side or ''
                ledger.save()
                report['ledgers'].append({'row': rec['row'], 'code': ledger.code, 'name': ledger.name,
                                          'kind': ledger.kind, 'action': 'created',
                                          'opening_balance': str(ledger.opening_balance)})
            else:
                changed = []
                if ledger.kind != kind:
                    raise ValueError(f'{ledger.code} {ledger.name} is a {ledger.get_kind_display()}, '
                                     f'not a {dict(LedgerAccount.KINDS)[kind]}')
                if opening is not None and (ledger.opening_balance != opening
                                            or (side and ledger.opening_side != side)):
                    ledger.opening_balance = opening
                    ledger.opening_side = side or ledger.opening_side
                    changed.append('opening balance')
                if rec['notes'] and ledger.notes != rec['notes']:
                    ledger.notes = rec['notes']
                    changed.append('notes')
                if changed:
                    ledger.save()
                report['ledgers'].append({'row': rec['row'], 'code': ledger.code, 'name': ledger.name,
                                          'kind': ledger.kind,
                                          'action': 'updated ' + ', '.join(changed) if changed else 'exists',
                                          'opening_balance': str(ledger.opening_balance)})
        except (ValueError, ValidationError) as exc:
            report['errors'].append(f'{where}: {exc}')


# ---------------------------------------------------------------------------
# Vouchers sheet
# ---------------------------------------------------------------------------

class _Ledgers:
    """Ledger lookup by code or name, loaded once per import."""

    def __init__(self):
        self.reload()

    def reload(self):
        self.by_code, self.by_name = {}, {}
        for a in LedgerAccount.objects.filter(is_active=True).select_related('customer', 'supplier'):
            self.by_code[a.code.lower()] = a
            self.by_name.setdefault(a.name.strip().lower(), []).append(a)

    def find(self, text):
        key = _text(text, 150).lower()
        if not key:
            raise ValueError('the ledger is blank')
        if key in self.by_code:
            return self.by_code[key]
        # "CU-0001 Mwananchi Traders" as the form shows it
        first = key.split()[0]
        if first in self.by_code and self.by_code[first].name.lower() == key[len(first):].strip():
            return self.by_code[first]
        hits = self.by_name.get(key, [])
        if len(hits) == 1:
            return hits[0]
        if len(hits) > 1:
            raise ValueError(f'"{text}" names {len(hits)} ledgers — use the code '
                             f'({", ".join(a.code for a in hits)})')
        raise ValueError(f'no ledger called "{text}" (use the code or the exact name from the Ledgers list)')


def _party_ledger(kind, name, ledgers, create):
    """The customer's or supplier's ledger for a party named on a row —
    optionally creating the party (and so the ledger) when it is new."""
    model = Customer if kind == 'customer' else Supplier
    party = model.objects.filter(name__iexact=name).first()
    if party is None:
        if not create:
            raise ValueError(f'no {kind} called "{name}" — add them first, or tick '
                             f'"create missing customers and suppliers"')
        party = model.objects.create(name=name)
        vouchers.ensure_ledger_for(party)
        ledgers.reload()
    ledger = getattr(party, 'ledger', None)
    if ledger is None:
        vouchers.ensure_ledger_for(party)
        party.refresh_from_db()
        ledger = party.ledger
        ledgers.reload()
    return party, ledger


def _parse_vouchers(rows):
    header, body = _header_and_rows(rows)
    if header is None:
        raise ImportError_('No header row found on the Vouchers sheet.')
    mapping = _map_headers(header, VOUCHER_ALIASES, ['type', 'date', 'ledger', 'side', 'amount'],
                           VOUCHER_HEADERS)
    if len(body) > MAX_ROWS:
        raise ImportError_(f'The sheet has more than {MAX_ROWS} rows — split it and upload in parts.')

    # Group the lines: one voucher per (ref, type, date). A blank ref is filled
    # from the invoice number, then from the row itself, so a single-row
    # file still groups sensibly.
    groups = OrderedDict()
    errors = []
    for row_no, raw in body:
        rec = {f: None for f in VOUCHER_ALIASES.values()}
        for index, field in mapping.items():
            if index < len(raw):
                rec[field] = raw[index]
        try:
            vtype = TYPE_ALIASES.get(_normalise(rec['type']))
            if not vtype:
                raise ValueError(f"type \"{_text(rec['type'])}\" is not one of: sales, purchase, receipt, "
                                 f"payment, contra, journal")
            when = parse_date(rec['date'])
            side = SIDE_ALIASES.get(_normalise(rec['side']))
            if not side:
                raise ValueError(f"side \"{_text(rec['side'])}\" must be Dr or Cr")
            amount = parse_amount(rec['amount'])
            if amount <= 0:
                raise ValueError('the amount must be more than nothing')
        except ValueError as exc:
            errors.append(f'Vouchers row {row_no}: {exc}')
            continue

        ref = _text(rec['ref'], 50) or _text(rec['invoice_number'], 50) or f'row {row_no}'
        key = (ref, vtype, when)
        g = groups.setdefault(key, {
            'ref': ref, 'type': vtype, 'date': when, 'rows': [], 'lines': [],
            'description': '', 'party': '', 'invoice_number': '', 'efd_number': '',
        })
        g['rows'].append(row_no)
        g['lines'].append({'row': row_no, 'ledger': rec['ledger'], 'side': side, 'amount': amount,
                           'narration': _text(rec['narration'], 200), 'against': _text(rec['against'], 50)})
        # Header fields: the first non-blank value in the group wins.
        for field in ('description', 'party', 'invoice_number', 'efd_number'):
            if not g[field] and _text(rec[field]):
                g[field] = _text(rec[field], 500 if field == 'description' else 50)
    return list(groups.values()), errors


# ---------------------------------------------------------------------------
# The run
# ---------------------------------------------------------------------------

class _Rollback(Exception):
    pass


def run(upload, user, commit=False, create_parties=False):
    """Read the file, apply the Ledgers sheet, post every voucher — all inside
    one transaction that is rolled back unless `commit` is true *and* nothing
    failed. Returns the report either way."""
    sheets = _read_sheets(upload)
    report = {'ok': False, 'committed': False, 'ledgers': [], 'vouchers': [], 'errors': [],
              'posted': 0, 'failed': 0, 'total': '0.00'}

    ledger_records = _parse_ledgers(sheets.get('ledgers') or []) if 'ledgers' in sheets else []
    groups, row_errors = _parse_vouchers(_pick_voucher_sheet(sheets))
    report['errors'].extend(row_errors)
    if not groups and not ledger_records:
        raise ImportError_('The file has no voucher rows and no ledgers to import.')

    try:
        with transaction.atomic():
            _apply_ledgers(ledger_records, report)
            ledgers = _Ledgers()
            total = Decimal('0.00')
            # Oldest first, so the voucher numbers run in date order.
            for g in sorted(groups, key=lambda g: (g['date'], g['rows'][0])):
                result = {'ref': g['ref'], 'type': g['type'], 'date': g['date'].isoformat(),
                          'rows': f"{g['rows'][0]}–{g['rows'][-1]}" if len(g['rows']) > 1 else str(g['rows'][0]),
                          'lines': len(g['lines']), 'party': g['party'], 'invoice_number': g['invoice_number'],
                          'debit': str(sum(l['amount'] for l in g['lines'] if l['side'] == 'debit')),
                          'credit': str(sum(l['amount'] for l in g['lines'] if l['side'] == 'credit')),
                          'status': 'ok', 'number': '', 'message': ''}
                try:
                    header = {'invoice_number': g['invoice_number'], 'efd_number': g['efd_number']}
                    party_ledger = None
                    if g['type'] in Voucher.INVOICE_TYPES and g['party']:
                        kind = 'customer' if g['type'] == 'sales' else 'supplier'
                        party, party_ledger = _party_ledger(kind, g['party'], ledgers, create_parties)
                        header[kind] = party.id
                    raw_lines = []
                    for line in g['lines']:
                        try:
                            # "Customer" / "Supplier" in the Ledger column means the
                            # party's own ledger, so the sheet need not spell its code.
                            if _normalise(line['ledger']) in PARTY_WORDS:
                                if party_ledger is None:
                                    raise ValueError(f'"{_text(line["ledger"])}" needs a Party on the row')
                                account = party_ledger
                            else:
                                account = ledgers.find(line['ledger'])
                        except ValueError as exc:
                            raise ValidationError({'lines': f"row {line['row']}: {exc}"})
                        raw = {'account': account.id, 'side': line['side'],
                               'amount': str(line['amount']), 'narration': line['narration']}
                        if line['against']:
                            raw['allocations'] = [_allocation_for(account, line)]
                        raw_lines.append(raw)
                    voucher = vouchers.post_voucher(g['type'], g['date'], g['description'], raw_lines,
                                                    user, header=header)
                    result['number'] = voucher.number
                    total += voucher.total
                    report['posted'] += 1
                except ValidationError as exc:
                    result['status'] = 'error'
                    result['message'] = _flatten(exc.detail)
                    report['failed'] += 1
                except ValueError as exc:
                    result['status'] = 'error'
                    result['message'] = str(exc)
                    report['failed'] += 1
                report['vouchers'].append(result)
            report['total'] = str(total)
            report['ok'] = not report['errors'] and report['failed'] == 0
            if not (commit and report['ok']):
                raise _Rollback()
            report['committed'] = True
    except _Rollback:
        # Voucher numbers handed out in the dry run are rolled back with it.
        for r in report['vouchers']:
            r['number'] = ''
    return report


def _allocation_for(account, line):
    """The `Against` column: set this line against the one open invoice or
    bill of the ledger's party whose reference matches — a till invoice
    number, PO-<n>, or the invoice number of a Sales/Purchase voucher."""
    if vouchers.allocation_side(account) != line['side']:
        raise ValidationError({'lines': (
            f"row {line['row']}: 'Against' only applies to a customer credited on a Receipt "
            f"or a supplier debited on a Payment")})
    wanted = line['against'].lower()
    rows = [r for r in vouchers.outstanding_for(account) if r['reference'].lower() == wanted]
    if not rows:
        raise ValidationError({'lines': (
            f"row {line['row']}: {account.name} has nothing outstanding called \"{line['against']}\"")})
    row = rows[0]
    amount = min(line['amount'], Decimal(row['outstanding']))
    return {row['target']: row['id'], 'amount': str(amount)}


def _flatten(detail):
    if isinstance(detail, dict):
        return '; '.join(_flatten(v) for v in detail.values())
    if isinstance(detail, (list, tuple)):
        return '; '.join(_flatten(v) for v in detail)
    return str(detail)


# ---------------------------------------------------------------------------
# The template
# ---------------------------------------------------------------------------

def build_template():
    """An .xlsx with the Vouchers and Ledgers sheets, worked examples, and the
    current chart of accounts to copy codes from."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    bold = Font(bold=True)
    fill = PatternFill('solid', fgColor='EEF1FB')
    wb = Workbook()

    ws = wb.active
    ws.title = 'Vouchers'
    ws.append(VOUCHER_HEADERS)
    example = [
        # A VAT credit sale: three rows, one Ref
        ['1', 'Sales', '2026-07-03', 'Customer', 'Dr', 1180000, 'Cement, invoice 0451', 'Kibo Traders',
         'INV-0451', 'EFD-88123', '', ''],
        ['1', 'Sales', '2026-07-03', 'Sales Revenue', 'Cr', 1000000, '', '', '', '', '', ''],
        ['1', 'Sales', '2026-07-03', 'Output VAT', 'Cr', 180000, '', '', '', '', '', ''],
        # A cash purchase
        ['2', 'Purchase', '2026-07-05', 'Purchases', 'Dr', 250000, 'Nails, supplier invoice 77', 'Mwenge Steel',
         'MS-77', '', '', ''],
        ['2', 'Purchase', '2026-07-05', 'Main Cash Book', 'Cr', 250000, '', '', '', '', '', ''],
        # The customer pays part of the invoice
        ['3', 'Receipt', '2026-07-20', 'CRDB Main', 'Dr', 500000, 'Part payment INV-0451', '', '', '', '', ''],
        ['3', 'Receipt', '2026-07-20', 'Kibo Traders', 'Cr', 500000, '', '', '', '', 'INV-0451', ''],
        # Rent paid from the bank
        ['4', 'Payment', '2026-07-31', 'Rent', 'Dr', 800000, 'July rent', '', '', '', '', ''],
        ['4', 'Payment', '2026-07-31', 'CRDB Main', 'Cr', 800000, '', '', '', '', '', ''],
    ]
    for row in example:
        ws.append(row)
    for cell in ws[1]:
        cell.font = bold
        cell.fill = fill
    for col, width in zip('ABCDEFGHIJKL', (8, 11, 12, 26, 6, 14, 30, 22, 14, 14, 14, 24)):
        ws.column_dimensions[col].width = width

    ls = wb.create_sheet('Ledgers')
    ls.append(LEDGER_HEADERS)
    ls.append(['', 'Equity Bank', 'bank', 2500000, 'Dr', 'Opening balance 1 July'])
    ls.append(['', 'Motor Vehicles', 'asset', 18000000, 'Dr', ''])
    ls.append(['', 'Capital', 'equity', 20500000, 'Cr', ''])
    for cell in ls[1]:
        cell.font = bold
        cell.fill = fill
    for col, width in zip('ABCDEF', (10, 28, 12, 16, 12, 30)):
        ls.column_dimensions[col].width = width

    cs = wb.create_sheet('Ledger List')
    cs.append(['Code', 'Name', 'Kind', 'Use on'])
    use = {
        'bank': 'Receipt Dr · Payment Cr · Contra · Sales Dr (cash sale) · Purchase Cr (cash purchase)',
        'cash': 'Receipt Dr · Payment Cr · Contra · Sales Dr (cash sale) · Purchase Cr (cash purchase)',
        'customer': 'Sales Dr (credit sale) · Receipt Cr',
        'supplier': 'Purchase Cr (credit purchase) · Payment Dr',
        'income': 'Sales Cr · Receipt Cr',
        'expense': 'Purchase Dr · Payment Dr',
        'asset': 'Purchase Dr · Journal',
        'liability': 'Journal · Payment Dr',
        'tax': 'Sales Cr (Output VAT) · Purchase Dr (Input VAT) · Payment Dr',
        'equity': 'Journal',
    }
    for a in LedgerAccount.objects.filter(is_active=True).order_by('kind', 'code'):
        cs.append([a.code, a.name, a.get_kind_display(), use.get(a.kind, '')])
    for cell in cs[1]:
        cell.font = bold
        cell.fill = fill
    for col, width in zip('ABCD', (10, 32, 16, 70)):
        cs.column_dimensions[col].width = width

    hs = wb.create_sheet('How to fill')
    for line in (
        'VOUCHERS sheet — one row per line of a voucher.',
        '  Ref: any text that groups the rows of one voucher (1, 2, 3… or the invoice number).',
        '  Type: Sales, Purchase, Receipt, Payment, Contra or Journal.',
        '  Date: YYYY-MM-DD or DD/MM/YYYY. Ledger: the code or exact name from the Ledger List sheet.',
        '  Side: Dr or Cr. Amount: a positive number. Debits and credits of a Ref must add up.',
        '  Party / Invoice No / EFD No: Sales and Purchase vouchers only — the customer or supplier, our '
        'invoice number (or theirs), and the EFD receipt number. Write "Customer" or "Supplier" as the Ledger '
        'on the party line and the party\'s own ledger is used.',
        '  Against: on a Receipt row crediting a customer (or a Payment row debiting a supplier), the '
        'invoice or bill number it settles: a till invoice number, PO-<n>, or the invoice number of a '
        'Sales/Purchase voucher. Left blank, the money is received on account.',
        '  Description applies to the voucher; Narration to one line.',
        'LEDGERS sheet — ledgers to create before the vouchers post, with opening balances. Customers and '
        'suppliers are not created here: name them on a voucher row.',
        'Each voucher type only accepts certain ledgers on each side — see the Use on column of Ledger List.',
        'Nothing posts unless the whole file is clean. "Check file" first; the report names the rows to fix.',
    ):
        hs.append([line])
    hs.column_dimensions['A'].width = 120
    return wb
