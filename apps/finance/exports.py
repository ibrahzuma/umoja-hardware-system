"""CSV and Excel export, shared by every accounting report.

One pair of functions behind one `export_response`, so every report exports
the same way and a new report gets both formats for free. `openpyxl` is
imported inside `export_xlsx` — the same lazy pattern the rest of the finance
app uses, so a CSV download never pays for the Excel machinery.
"""

import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone


def _cell(value):
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, 'pk') and not isinstance(value, (str, int)):
        return str(value)
    return value


def export_csv(filename, columns, rows):
    """CSV with a UTF-8 BOM, so Excel opens it with the right encoding — the
    same courtesy `static/js/table_export.js` does on screen."""
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    response.write('﻿')
    writer = csv.writer(response)
    writer.writerow(columns)
    for row in rows:
        writer.writerow([
            '' if value is None
            else (f"{value:.2f}" if isinstance(value, Decimal) else str(value))
            for value in row
        ])
    return response


def export_xlsx(filename, columns, rows, title=''):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (title or filename)[:30] or 'Report'

    row_index = 1
    if title:
        sheet.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
        sheet.cell(row=2, column=1, value=f"Generated {timezone.localtime():%d/%m/%Y %H:%M}")
        row_index = 4

    header_fill = PatternFill('solid', fgColor='DDE3EA')
    for column, name in enumerate(columns, start=1):
        cell = sheet.cell(row=row_index, column=column, value=name)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        row_index += 1
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column, value=_cell(value))
            if isinstance(value, Decimal):
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal='right')
            elif isinstance(value, (date, datetime)):
                cell.number_format = 'DD/MM/YYYY'

    for column in range(1, len(columns) + 1):
        widest = max(
            (len(str(row[column - 1])) if column - 1 < len(row) and row[column - 1] is not None else 0)
            for row in [columns] + list(rows)
        )
        sheet.column_dimensions[get_column_letter(column)].width = min(max(12, widest + 2), 60)

    buffer = BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def export_response(fmt, filename, columns, rows, title=''):
    rows = list(rows)
    if fmt == 'xlsx':
        return export_xlsx(filename, columns, rows, title=title)
    return export_csv(filename, columns, rows)
