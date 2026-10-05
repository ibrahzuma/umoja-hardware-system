"""Bulk upload of the chart of accounts from an Excel (.xlsx) or CSV file.

Columns, matched case-insensitively and in any order:

    Code | Name | Type | Parent Code | Kind | Currency | Opening Balance | Dr/Cr | Description

    Type     Asset / Liability / Equity / Income / Expense
    Kind     blank (ordinary ledger), Group, Bank, Cash, Customer, Supplier,
             Income, Expense, Asset, Liability, Tax, Equity,
             Output VAT, Input VAT, Customer Control, Supplier Control
    Currency a currency code (blank = base currency). A ledger with a
             currency can only be used on vouchers in that currency.
    Dr/Cr    which side the opening balance sits on; defaults to the natural
             side of the account type.

Rows are read in file order, so a group can be listed above the ledgers that
sit under it. An existing code *updates* that ledger rather than failing —
which is what makes the upload usable twice, to correct a mistake — except
that a ledger with entries in the General Ledger keeps its type, because
moving a posted ledger between asset and expense would silently rewrite the
statements.
"""

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Optional

from django.core.exceptions import ValidationError
from django.db import transaction

from .models import Currency, LedgerAccount
from .money import ZERO, quantize

HEADERS = ['Code', 'Name', 'Type', 'Parent Code', 'Kind', 'Currency',
           'Opening Balance', 'Dr/Cr', 'Description']
ALIASES = {
    'code': ('code', 'account code', 'ledger code', 'account no', 'no'),
    'name': ('name', 'account name', 'ledger name', 'ledger', 'account'),
    'type': ('type', 'account type', 'class'),
    'parent': ('parent code', 'parent', 'parent account', 'group code', 'group', 'under'),
    'kind': ('kind', 'ledger kind', 'special', 'flag', 'nature'),
    'currency': ('currency', 'currency code', 'ccy'),
    'opening': ('opening balance', 'opening', 'balance', 'opening bal'),
    'side': ('dr/cr', 'dr cr', 'side', 'balance type', 'opening balance type', 'dr or cr'),
    'description': ('description', 'notes', 'remarks'),
}
REQUIRED = ('code', 'name', 'type')

TYPES = {
    'asset': LedgerAccount.ASSET, 'assets': LedgerAccount.ASSET, 'a': LedgerAccount.ASSET,
    'liability': LedgerAccount.LIABILITY, 'liabilities': LedgerAccount.LIABILITY,
    'l': LedgerAccount.LIABILITY,
    'equity': LedgerAccount.EQUITY, 'capital': LedgerAccount.EQUITY, 'e': LedgerAccount.EQUITY,
    'income': LedgerAccount.INCOME, 'revenue': LedgerAccount.INCOME,
    'sales': LedgerAccount.INCOME, 'i': LedgerAccount.INCOME,
    'expense': LedgerAccount.EXPENSE, 'expenses': LedgerAccount.EXPENSE,
    'x': LedgerAccount.EXPENSE,
}

# The Kind column decides both the ledger's `kind` (which voucher dropdown it
# appears in) and the flags on it. `None` for kind means "work it out from the
# account type", which is what a blank Kind does.
KINDS = {
    '': (None, {}), 'ledger': (None, {}), 'normal': (None, {}), 'none': (None, {}),
    'group': (None, {'is_group': True}),
    'header': (None, {'is_group': True}),
    'bank': ('bank', {}), 'bank account': ('bank', {}),
    'cash': ('cash', {}), 'cash book': ('cash', {}), 'cashbook': ('cash', {}),
    'customer': ('customer', {}), 'supplier': ('supplier', {}),
    'income': ('income', {}), 'expense': ('expense', {}),
    'asset': ('asset', {}), 'liability': ('liability', {}),
    'tax': ('tax', {}), 'equity': ('equity', {}),
    'output vat': ('tax', {'vat_kind': LedgerAccount.VAT_OUTPUT}),
    'vat output': ('tax', {'vat_kind': LedgerAccount.VAT_OUTPUT}),
    'vat on sales': ('tax', {'vat_kind': LedgerAccount.VAT_OUTPUT}),
    'sales vat': ('tax', {'vat_kind': LedgerAccount.VAT_OUTPUT}),
    'input vat': ('tax', {'vat_kind': LedgerAccount.VAT_INPUT}),
    'vat input': ('tax', {'vat_kind': LedgerAccount.VAT_INPUT}),
    'vat on purchases': ('tax', {'vat_kind': LedgerAccount.VAT_INPUT}),
    'purchase vat': ('tax', {'vat_kind': LedgerAccount.VAT_INPUT}),
    'customer control': ('customer', {'is_customer_control': True}),
    'customers control': ('customer', {'is_customer_control': True}),
    'receivables control': ('customer', {'is_customer_control': True}),
    'debtors control': ('customer', {'is_customer_control': True}),
    'accounts receivable': ('customer', {'is_customer_control': True}),
    'supplier control': ('supplier', {'is_supplier_control': True}),
    'suppliers control': ('supplier', {'is_supplier_control': True}),
    'payables control': ('supplier', {'is_supplier_control': True}),
    'creditors control': ('supplier', {'is_supplier_control': True}),
    'accounts payable': ('supplier', {'is_supplier_control': True}),
}
KIND_LABELS = ['(blank) — ordinary ledger', 'Group', 'Bank', 'Cash', 'Customer', 'Supplier',
               'Income', 'Expense', 'Asset', 'Liability', 'Tax', 'Equity',
               'Output VAT', 'Input VAT', 'Customer Control', 'Supplier Control']

# Which `kind` a blank Kind column implies, from the account type alone.
KIND_FOR_TYPE = {
    LedgerAccount.ASSET: 'asset', LedgerAccount.LIABILITY: 'liability',
    LedgerAccount.EQUITY: 'equity', LedgerAccount.INCOME: 'income',
    LedgerAccount.EXPENSE: 'expense',
}
CATEGORY_FOR_TYPE = {
    LedgerAccount.ASSET: 'CURRENT_ASSET', LedgerAccount.LIABILITY: 'CURRENT_LIABILITY',
    LedgerAccount.EQUITY: 'EQUITY', LedgerAccount.INCOME: 'REVENUE',
    LedgerAccount.EXPENSE: 'OPERATING_EXPENSE',
}

EXAMPLES = [
    ['1000', 'CURRENT ASSETS', 'Asset', '', 'Group', '', '', '', ''],
    ['1001', 'MAIN CASH BOOK', 'Asset', '1000', 'Cash', '', '250000', 'Dr', 'Main till'],
    ['1050', 'NMB CURRENT ACCOUNT', 'Asset', '1000', 'Bank', '', '1500000', 'Dr', ''],
    ['1052', 'NMB USD ACCOUNT', 'Asset', '1000', 'Bank', 'USD', '', '', 'Foreign currency account'],
    ['1100', 'ACCOUNTS RECEIVABLE', 'Asset', '', 'Customer Control', '', '', '', ''],
    ['4001', 'STANDARD RATED SALES', 'Income', '', 'Income', '', '', '', ''],
    ['4002', 'EXEMPT SALES', 'Income', '', 'Income', '', '', '', ''],
    ['5001', 'PURCHASES', 'Expense', '', 'Expense', '', '', '', ''],
    ['2200', 'ACCOUNTS PAYABLE', 'Liability', '', 'Supplier Control', '', '', '', ''],
    ['8000', 'VAT ACCOUNT', 'Liability', '', 'Group', '', '', '', ''],
    ['8001', 'VAT 18% — ON PURCHASES', 'Liability', '8000', 'Input VAT', '', '', '', ''],
    ['8002', 'VAT 18% — ON SALES', 'Liability', '8000', 'Output VAT', '', '', '', ''],
]


@dataclass
class RowResult:
    row: int
    account: Optional[LedgerAccount] = None
    created: bool = False
    errors: list = field(default_factory=list)

    @property
    def ok(self):
        return not self.errors


class AccountImportError(Exception):
    pass


def build_template():
    """The workbook to fill in: the headings, worked examples, the rules, and
    the chart as it stands so codes can be copied rather than retyped."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Ledgers'
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0D6EFD')
    for row in EXAMPLES:
        sheet.append(row)
    for index, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(header) + 6)
    sheet.column_dimensions['B'].width = 34

    notes = workbook.create_sheet('Instructions')
    for line in [
        "One row per ledger. Delete the example rows before uploading — or edit them, since an "
        "existing Code updates that ledger rather than failing.",
        "Code: a unique ledger code, e.g. 1050.   Name: what the ledger is called.",
        "Type: Asset, Liability, Equity, Income or Expense.",
        "Parent Code: the code of the group (or customer / supplier control account) this ledger "
        "sits under — optional. List a group above its children.",
        "Kind: " + ', '.join(KIND_LABELS) + ".",
        "Currency: a currency code set up under Currencies — blank means the base currency. A "
        "ledger with a currency can only be used on vouchers in that currency.",
        "Opening Balance: optional. Dr/Cr says which side it sits on, and defaults to Dr for "
        "assets and expenses, Cr otherwise. Enter it as a positive figure either way.",
        "A ledger that already has General Ledger entries keeps its Type — changing it would "
        "rewrite the statements behind everyone's back.",
    ]:
        notes.append([line])
    notes.column_dimensions['A'].width = 120

    existing = workbook.create_sheet('Existing ledgers')
    existing.append(['Code', 'Name', 'Type', 'Parent', 'Kind', 'Currency'])
    for cell in existing[1]:
        cell.font = Font(bold=True)
    for account in LedgerAccount.objects.select_related('parent', 'currency').order_by('code'):
        existing.append([account.code, account.name, account.get_account_type_display() or '',
                         account.parent.code if account.parent_id else '',
                         account.ledger_kind,
                         account.currency.code if account.currency_id else ''])
    for column, width in (('A', 14), ('B', 40), ('C', 14), ('D', 12), ('E', 20), ('F', 10)):
        existing.column_dimensions[column].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------- parsing

def _norm(value):
    return ' '.join(str(value or '').strip().lower().split())


def _text(value):
    text = str(value if value is not None else '').strip()
    if text.endswith('.0') and text[:-2].isdigit():
        text = text[:-2]       # Excel stored the code as a number
    return text


def read_rows(uploaded):
    name = (getattr(uploaded, 'name', '') or '').lower()
    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        from openpyxl import load_workbook
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        rows = workbook.worksheets[0].iter_rows(values_only=True)
    elif name.endswith('.csv') or name.endswith('.txt'):
        text = uploaded.read().decode('utf-8-sig', errors='replace')
        rows = iter(csv.reader(io.StringIO(text)))
    else:
        raise AccountImportError("Upload an Excel (.xlsx) or a CSV file.")

    try:
        headers = [_norm(h) for h in next(rows)]
    except StopIteration:
        raise AccountImportError("The file is empty.")

    mapping = {}
    for key, aliases in ALIASES.items():
        for header in headers:
            if header in aliases:
                mapping[key] = headers.index(header)
                break
    missing = [key.title() for key in REQUIRED if key not in mapping]
    if missing:
        raise AccountImportError("Missing column(s): " + ', '.join(missing))

    out = []
    for index, values in enumerate(rows, start=2):
        values = list(values or [])
        if not any(str(v).strip() for v in values if v is not None):
            continue

        def get(key, values=values):
            position = mapping.get(key)
            return values[position] if position is not None and position < len(values) else None

        out.append((index, {key: get(key) for key in ALIASES}))
    if not out:
        raise AccountImportError("No data rows below the headings.")
    return out


# ---------------------------------------------------------------------- import

class AccountImportService:
    def __init__(self, request=None, user=None):
        self.request = request
        self.user = user

    def run(self, uploaded):
        return [self._import_row(row_no, data) for row_no, data in read_rows(uploaded)]

    def _import_row(self, row_no, data):
        result = RowResult(row=row_no)
        try:
            parsed = self._parse(data)
        except ValueError as exc:
            result.errors.append(str(exc))
            return result

        account = LedgerAccount.objects.filter(code__iexact=parsed['code']).first()
        created = account is None
        if created:
            account = LedgerAccount(code=parsed['code'])
        elif (account.gl_entries.exists()
              and account.account_type != parsed['account_type']):
            result.errors.append(
                f"{parsed['code']} already has General Ledger entries, so its Type cannot "
                f"be changed.")
            return result

        if parsed['parent'] is not None and not created and parsed['parent'].pk == account.pk:
            result.errors.append("A ledger cannot be its own parent.")
            return result

        account.name = parsed['name']
        account.account_type = parsed['account_type']
        account.kind = parsed['kind']
        account.parent = parsed['parent']
        account.currency = parsed['currency']
        account.is_active = True
        # Every flag is set explicitly, so a re-upload that drops a flag
        # actually clears it rather than leaving it stuck on.
        account.is_group = parsed['flags'].get('is_group', False)
        account.is_customer_control = parsed['flags'].get('is_customer_control', False)
        account.is_supplier_control = parsed['flags'].get('is_supplier_control', False)
        account.vat_kind = parsed['flags'].get('vat_kind', '')
        if parsed['has_opening']:
            account.opening_balance = parsed['opening']
            account.opening_side = parsed['side']
        if parsed['description']:
            account.notes = parsed['description']
        if created and not account.category:
            account.category = CATEGORY_FOR_TYPE.get(parsed['account_type'], '')

        try:
            with transaction.atomic():
                account.full_clean(exclude=['code'])
                account.save()
        except ValidationError as exc:
            for field_name, messages in exc.message_dict.items():
                prefix = '' if field_name == '__all__' else f"{field_name}: "
                result.errors.extend(prefix + m for m in messages)
            return result
        result.account, result.created = account, created
        return result

    def _parse(self, data):
        code, name = _text(data['code']), _text(data['name'])
        if not code:
            raise ValueError("Code is required.")
        if not name:
            raise ValueError("Name is required.")

        type_key = _norm(data['type'])
        if type_key not in TYPES:
            raise ValueError(f"Unknown Type '{_text(data['type'])}' — use Asset, Liability, "
                             f"Equity, Income or Expense.")
        account_type = TYPES[type_key]

        kind_key = _norm(data['kind'])
        if kind_key not in KINDS:
            raise ValueError(f"Unknown Kind '{_text(data['kind'])}' — use "
                             + ', '.join(KIND_LABELS) + ".")
        kind, flags = KINDS[kind_key]
        if kind is None:
            kind = KIND_FOR_TYPE[account_type]

        currency = None
        currency_code = _text(data['currency']).upper()
        if currency_code:
            currency = Currency.objects.filter(code__iexact=currency_code, is_active=True).first()
            if currency is None:
                raise ValueError(f"Currency '{currency_code}' is not set up under Currencies.")

        parent = None
        parent_code = _text(data['parent'])
        if parent_code:
            parent = LedgerAccount.objects.filter(code__iexact=parent_code).first()
            if parent is None:
                raise ValueError(f"Parent Code '{parent_code}' does not exist — list a group "
                                 f"above the ledgers under it.")

        has_opening = data['opening'] is not None and _text(data['opening']) != ''
        opening = self._amount(data['opening']) if has_opening else ZERO

        side_key = _norm(data['side'])
        if side_key in ('dr', 'debit', 'd'):
            side = 'debit'
        elif side_key in ('cr', 'credit', 'c'):
            side = 'credit'
        elif side_key == '':
            side = ('debit' if account_type in (LedgerAccount.ASSET, LedgerAccount.EXPENSE)
                    else 'credit')
        else:
            raise ValueError(f"Dr/Cr has to be Dr or Cr, not '{_text(data['side'])}'.")

        return {'code': code, 'name': name, 'account_type': account_type, 'kind': kind,
                'flags': flags, 'currency': currency, 'parent': parent,
                'has_opening': has_opening, 'opening': opening, 'side': side,
                'description': _text(data['description'])}

    @staticmethod
    def _amount(value):
        text = str(value).replace(',', '').replace(' ', '').strip()
        try:
            amount = quantize(Decimal(text))
        except (InvalidOperation, ValueError):
            raise ValueError(f"'{value}' is not a readable Opening Balance.")
        if amount < 0:
            raise ValueError("An Opening Balance cannot be negative — use the Dr/Cr column "
                             "to say which side it sits on.")
        return amount
