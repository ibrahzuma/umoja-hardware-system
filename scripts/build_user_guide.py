"""Generate docs/Umoja_User_Guide.pptx — a role-by-role layman guide to the
Umoja Hardware System, plus an end-to-end sale walkthrough and a mobile
companion app section.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "screens"
OUT = ROOT / "docs" / "Umoja_User_Guide.pptx"

BRAND = RGBColor(0x40, 0x55, 0xE6)
INK = RGBColor(0x14, 0x1A, 0x23)
MUTED = RGBColor(0x6B, 0x76, 0x86)
ACCENT = RGBColor(0x16, 0xA3, 0x7B)
DANGER = RGBColor(0xD9, 0x3B, 0x3B)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


@dataclass
class Step:
    title: str
    body: str
    screen: str | None = None


@dataclass
class Chapter:
    code: str
    name: str
    who: str
    daily_summary: str
    steps: list[Step] = field(default_factory=list)
    tips: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

ROLES: list[Chapter] = [
    Chapter(
        code="admin",
        name="System Administrator",
        who="The owner or IT lead who sets the system up and keeps it healthy.",
        daily_summary=(
            "You configure the business once (branches, categories, users, roles) "
            "and then mostly use the dashboard to spot problems. You can do "
            "everything any other role can do."
        ),
        steps=[
            Step("Sign in",
                 "Open https://umoja.ehub.co.tz, type your username and password, click Sign in.",
                 "01_login.png"),
            Step("Admin dashboard",
                 "After login you land on the Dashboard. The cards at top tell "
                 "you what needs attention right now: orders ready to dispatch, "
                 "low-stock items, pending approvals, today's revenue.",
                 "role_admin_dashboard.png"),
            Step("Create branches",
                 "Inventory > Branches > Add branch. Every product stock and "
                 "every sale is tied to a branch — set yours up before adding "
                 "any other data.",
                 "03_branches.png"),
            Step("Create categories",
                 "Inventory > Categories. Each category has a commission % that "
                 "is applied to every sale of products in that category. "
                 "Historical sales keep their old rate even if you change it later.",
                 "04_categories.png"),
            Step("Manage users",
                 "Users. Each user has a role that decides which menu items "
                 "they see and which actions they can take.",
                 "role_admin_users_list.png"),
            Step("Company settings",
                 "Settings. Company name, default currency (TZS), and default "
                 "tax rate. These show up on receipts and invoices.",
                 "role_admin_settings.png"),
        ],
        tips=[
            "Run `python manage.py create_roles` once after install to seed all role groups.",
            "Keep the SECRET_KEY out of git — set it in /var/www/app/.env on the server.",
            "Reports menu shows aging, profitability by category, and commissions per sales rep.",
        ],
    ),
    Chapter(
        code="store_manager",
        name="Store Manager",
        who="The supervisor on the floor who approves orders and runs dispatch.",
        daily_summary=(
            "Sales reps create orders all day. Your job is to approve or "
            "decline them, then dispatch the approved ones — assigning a "
            "store keeper to pick the goods and a vehicle to deliver them."
        ),
        steps=[
            Step("Your dashboard",
                 "The Store Manager dashboard surfaces a Manager View — "
                 "ready-to-dispatch count, low-stock count, active fleet.",
                 "role_store_manager_dashboard.png"),
            Step("Order Management — approve or decline",
                 "Sales > Order management. Each pending row has Approve "
                 "and Decline. Approve only when you've checked the "
                 "customer is real and the price is right.",
                 "role_store_manager_order_management.png"),
            Step("Dispatch Board — assign & ship",
                 "Sales > Dispatch board. Pick an approved order, choose "
                 "a Store Keeper to pick the goods, pick a Vehicle (or "
                 "type hired transport in Lorry info), click Dispatch.",
                 "role_store_manager_dispatch_board.png"),
            Step("Track delivery vehicles",
                 "Sales > Vehicles. See which vehicles are Active vs busy "
                 "In transit. Mark a returned vehicle Active again so it "
                 "can be assigned to the next dispatch.",
                 "role_store_manager_vehicles.png"),
            Step("Move stock between branches",
                 "Inventory > Transfer. Use when one branch is low and "
                 "another is overstocked. Stock moves out of source, into destination.",
                 "role_store_manager_stock_transfer.png"),
        ],
        tips=[
            "Stock only moves at Dispatch — Approve does not yet deduct stock.",
            "If a vehicle stays 'busy' after the driver came back, click Return to free it up.",
            "Decline within minutes — the customer is usually still on the line.",
        ],
    ),
    Chapter(
        code="sales_rep",
        name="Sales Representative",
        who="The cashier or salesperson at the counter or on the phone.",
        daily_summary=(
            "You take orders all day. You enter products and customer, record "
            "payment (cash, mobile money, bank transfer, or credit), and submit. "
            "The order then waits for your manager to approve and dispatch it."
        ),
        steps=[
            Step("Your dashboard",
                 "Sales Rep dashboards focus on Commerce — your total orders, "
                 "low stock items, and a Sales (Shop) workspace.",
                 "role_sales_rep_dashboard.png"),
            Step("Open the POS",
                 "Sales > POS. Search products, add to cart, pick a customer, "
                 "record payment.",
                 "role_sales_rep_pos.png"),
            Step("Sales list",
                 "Sales > Sales. Your past sales with status (Pending, Approved, "
                 "Dispatched, Cancelled) and the amount, paid vs balance.",
                 "role_sales_rep_sales_list.png"),
            Step("Customers",
                 "Customers. Track existing customers and add new ones. Each "
                 "customer's sales history is one click away.",
                 "role_sales_rep_customers.png"),
            Step("Quotations",
                 "Quotations. Use for written price estimates — does NOT touch "
                 "stock. Convert into a Sale with one click when the customer commits.",
                 "role_sales_rep_quotations.png"),
        ],
        tips=[
            "Commission is calculated automatically based on the product's category.",
            "If a product isn't appearing in search, it might be out of stock at your branch.",
            "Use Quotations for big-ticket items the customer wants to think about.",
        ],
    ),
    Chapter(
        code="procurement",
        name="Procurement Officer (Afisa Ugavi)",
        who="The person responsible for ordering goods from suppliers.",
        daily_summary=(
            "You manage suppliers and raise Purchase Orders. Once goods "
            "arrive, the Stock Controller receives them against your PO."
        ),
        steps=[
            Step("Your dashboard",
                 "Procurement dashboards are scoped to Afisa Ugavi tasks.",
                 "role_procurement_dashboard.png"),
            Step("Manage suppliers",
                 "Inventory > Suppliers. Add new suppliers with their contact "
                 "info. Each Purchase Order is tied to one supplier.",
                 "role_procurement_suppliers.png"),
            Step("Purchase Orders",
                 "Inventory > Purchase orders. The PO list shows open, "
                 "received, and closed POs.",
                 "role_procurement_purchase_orders.png"),
            Step("Raise a new PO",
                 "Pick the supplier, add the line items (product + quantity + "
                 "agreed price). The PO is tracked in the system until goods arrive.",
                 "role_procurement_po_create.png"),
            Step("Inbound fleet",
                 "Inventory > Trucks (for inbound deliveries from suppliers). "
                 "These are SEPARATE from the outbound Vehicles used by Sales.",
                 "role_procurement_trucks.png"),
        ],
        tips=[
            "Trucks (this section) are for INBOUND from suppliers. Vehicles (Sales section) are for OUTBOUND to customers — don't mix them.",
            "Always raise the PO before the supplier ships — the Stock Controller's GRN matches against it line by line.",
        ],
    ),
    Chapter(
        code="stock_controller",
        name="Stock Controller",
        who="The receiving clerk who books incoming goods into the system.",
        daily_summary=(
            "When a delivery arrives from a supplier, you receive it against "
            "the PO with a Goods Received Note. That's what actually increases "
            "stock levels in the system."
        ),
        steps=[
            Step("Your dashboard",
                 "Stock Controller dashboards focus on receiving and stock health.",
                 "role_stock_controller_dashboard.png"),
            Step("GRN list",
                 "Inventory > GRNs. See past Goods Received Notes — one per "
                 "delivery received.",
                 "role_stock_controller_grns.png"),
            Step("Receive goods (new GRN)",
                 "Inventory > GRNs > Create. Pick the matching PO, then enter "
                 "the quantities ACTUALLY received per product. Stock is added "
                 "to your branch the moment you save.",
                 "role_stock_controller_grn_create.png"),
            Step("Live stock per branch",
                 "Inventory > Inventory. See current quantities per product per "
                 "branch. A red badge means the item is at or below the low-stock "
                 "threshold.",
                 "role_stock_controller_inventory.png"),
        ],
        tips=[
            "If the GRN quantity is less than the PO quantity, that's normal — the PO stays open for the rest.",
            "Set sensible low-stock thresholds on each Stock row so the dashboard alerts trigger before you actually run out.",
        ],
    ),
    Chapter(
        code="store_keeper",
        name="Store Keeper",
        who="The warehouse staff who picks goods for dispatch.",
        daily_summary=(
            "When the Store Manager assigns you a sale on the dispatch board, "
            "you go pull the goods, hand them to the driver, and the system "
            "records that the order has left."
        ),
        steps=[
            Step("Your dashboard",
                 "Store Keeper sees a minimal dashboard — just the stock and "
                 "sales information needed to pick orders.",
                 "role_store_keeper_dashboard.png"),
            Step("See assignments",
                 "Sales > Dispatch. Filter to orders assigned to you. Each row "
                 "shows the customer, the products to pick, and the delivery vehicle.",
                 "role_store_keeper_dispatch.png"),
            Step("Sales list",
                 "Sales > Sales. Cross-check the sale before picking — confirm "
                 "the customer, the products, and the quantities.",
                 "role_store_keeper_sales_list.png"),
        ],
        tips=[
            "If something on the order is short, raise it with the Store Manager BEFORE dispatch — once dispatched the stock is decremented.",
            "Print the delivery note from the sale's actions menu — it has the customer's address.",
        ],
    ),
    Chapter(
        code="accountant",
        name="Accountant",
        who="The bookkeeper who tracks money in and money out.",
        daily_summary=(
            "You log expenses, supplier payments, tax payments, and other "
            "income. The system already tracks sales transactions automatically "
            "— you mostly handle the non-sale side."
        ),
        steps=[
            Step("Your dashboard",
                 "Accountant dashboards focus on cash flow — expenses, income, taxes.",
                 "role_accountant_dashboard.png"),
            Step("Expenses list",
                 "Finance > Expenses. See the running list of recorded expenses, "
                 "filter by category or date.",
                 "role_accountant_expenses.png"),
            Step("Record an expense",
                 "Finance > Expenses > Add expense. Pick a category, type the "
                 "amount, attach the receipt photo. The expense lands in the "
                 "current month's report straight away.",
                 "role_accountant_expense_create.png"),
            Step("Other income",
                 "Finance > Income. Use for income that ISN'T a product sale "
                 "(rent received, service income, etc.).",
                 "role_accountant_income.png"),
            Step("Tax payments",
                 "Finance > Taxes. Pick the tax type (VAT, PAYE, SDL, …), "
                 "the period, and the amount paid. Keep the TRA receipt "
                 "number in the reference field.",
                 "role_accountant_taxes.png"),
            Step("Reports",
                 "Reports. Aging, Profitability by category, Commissions per "
                 "sales rep, Debtors.",
                 "role_accountant_reports.png"),
        ],
        tips=[
            "Always attach a receipt image to expenses — Tanzanian audits are picky and the system stores them with the row.",
            "The Debtors report uses Sale.total_amount minus sum of Transactions — make sure cash payments are entered, otherwise the customer looks like a debtor.",
        ],
    ),
]


WORKFLOW = Chapter(
    code="workflow",
    name="End-to-end sale: from POS to delivery",
    who="A real sale captured live on the system, showing each role's view at every transition.",
    daily_summary=(
        "Follow one Walk-in Customer's order from the moment the Sales "
        "Rep punches it in until the truck rolls out. The screenshots "
        "below are real captures from each user's logged-in session on "
        "the live site."
    ),
    steps=[
        Step("1. Sales Rep — open POS",
             "The Sales Rep logs in and opens Sales > POS. They search products, "
             "set quantity and price, pick the customer, take payment. When they "
             "Submit, the sale lands in 'Pending approval' state.",
             "wf_01_salesrep_pos.png"),
        Step("2. Sales Rep — sale is pending",
             "The new sale shows up in the rep's Sales list with status PENDING. "
             "Stock has NOT moved yet. The rep can take the next customer right "
             "away — approval is the manager's job.",
             "wf_02_salesrep_pending.png"),
        Step("3. Store Manager — sees pending order",
             "The Store Manager opens Sales > Order management. The new order "
             "is listed with Approve / Decline buttons. They check it's a real "
             "customer and the price is right.",
             "wf_03_smgr_orders_pending.png"),
        Step("4. Store Manager — approves",
             "After clicking Approve, the order disappears from the Pending list. "
             "Stock is still untouched at this point — Approve only means the "
             "order is cleared for dispatch.",
             "wf_04_smgr_orders_approved.png"),
        Step("5. Store Manager — opens Dispatch Board",
             "The approved order is now on Sales > Dispatch. The manager picks "
             "the Store Keeper who will pull the goods, picks a Vehicle (or "
             "types hired lorry info), and clicks Dispatch.",
             "wf_05_smgr_dispatch_before.png"),
        Step("6. Store Manager — order dispatched",
             "Now stock IS deducted. The order moves to DISPATCHED status, the "
             "vehicle is marked In Transit, and PDF invoice + delivery note are "
             "generated. The driver leaves with the goods.",
             "wf_06_smgr_dispatch_after.png"),
        Step("7. Store Manager — sales list",
             "Cross-checking the sale list now shows the order with Dispatched "
             "status, the picker assigned, and the lorry info.",
             "wf_07_smgr_sales_dispatched.png"),
        Step("8. Store Keeper — picks and confirms",
             "The Store Keeper sees the dispatched sale in their feed — they "
             "pull the goods physically and load them on the truck shown.",
             "wf_08_keeper_sales.png"),
        Step("9. Accountant — record the expense behind the sale",
             "Once the truck has rolled, the Accountant logs the fuel / freight "
             "expense via Finance > Expenses > Add expense. Receipt photo attached.",
             "wf_09_accountant_expense_form.png"),
        Step("10. Accountant — expenses list",
             "The new expense shows up in the running list, categorized, with "
             "the receipt image one click away.",
             "wf_10_accountant_expenses_list.png"),
    ],
    tips=[
        "Stock moves at step 6 (Dispatch), not at step 4 (Approve).",
        "Commission is frozen at step 1 (sale creation) based on the product's category — changing the category rate later doesn't affect historical sales.",
        "If the sale is later deleted, stock is restored automatically.",
    ],
)


MOBILE = Chapter(
    code="mobile",
    name="Mobile companion app (Umoja Pro)",
    who="Sales reps and field staff who work away from a desktop.",
    daily_summary=(
        "Umoja Pro is the Flutter mobile app for Android phones. It signs "
        "in with the same username/password as the web site and uses the "
        "same data — the four tabs cover the daily field-sales tasks."
    ),
    steps=[
        Step("Sign in",
             "Open Umoja Pro from the app drawer. The icon is the Umoja "
             "Hardware logo. Type your username and password and tap "
             "Sign in. The app remembers your token so you stay signed "
             "in across launches — sign out from the profile menu in the app bar.",
             "mobile_01_login.png"),
        Step("Four tabs at the bottom",
             "After login the app shows a four-tab bottom bar:\n\n"
             "  • POS — open a new sale, search products, take payment.\n"
             "  • Stock — see current stock per branch, with low-stock alerts.\n"
             "  • Quotes — list of quotations; tap to convert to a sale.\n"
             "  • Customers — directory; add new customers; see a customer's sales history.\n\n"
             "Each tab pulls the same data as the website over HTTPS to "
             "umoja.ehub.co.tz.",
             None),
        Step("Talks to the same backend",
             "The app calls /api-token-auth/ once for the token, then uses "
             "every other /api/ endpoint with that token in the Authorization "
             "header. Sales created on the phone show up in the Manager's "
             "Order Management screen on the website in real time, and stock "
             "deductions broadcast over WebSocket so other devices update "
             "without a refresh.",
             None),
    ],
    tips=[
        "If the Sign-in button does nothing, you're probably on the office Wi-Fi but the API URL is the local default. Reinstall a build made with --dart-define=API_BASE_URL=https://umoja.ehub.co.tz.",
        "Pull-to-refresh on each tab to force a re-fetch.",
        "Sign out from the profile menu in the app bar; this clears the token from secure storage.",
    ],
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def set_size(prs: Presentation) -> None:
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H


def add_textbox(slide, left, top, width, height, *, text: str,
                size: int = 18, bold: bool = False,
                color: RGBColor = INK, align=PP_ALIGN.LEFT) -> None:
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Calibri"


def add_band(slide, top, height, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), top, SLIDE_W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def cover(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_band(slide, Emu(0), SLIDE_H, INK)
    add_textbox(slide, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.2),
                text="UMOJA HARDWARE SYSTEM",
                size=44, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    add_textbox(slide, Inches(0.8), Inches(3.5), Inches(11.7), Inches(0.8),
                text="A role-by-role guide, with a real end-to-end sale walkthrough",
                size=20, color=RGBColor(0xB9, 0xC4, 0xFF))
    add_textbox(slide, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.6),
                text="umoja.ehub.co.tz   ·   Umoja Pro mobile app", size=14,
                color=RGBColor(0x9A, 0xA4, 0xB5))


def intro(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(0.6), Inches(0.5), Inches(12), Inches(0.8),
                text="How to read this guide", size=32, bold=True, color=INK)
    add_textbox(slide, Inches(0.6), Inches(1.5), Inches(12), Inches(1.3),
                text=("This guide has three parts:\n\n"
                      "1.  Each role gets its own chapter — find your role and read it end to end.\n"
                      "2.  The 'End-to-end sale' chapter shows one real sale going through every role on the live system.\n"
                      "3.  The 'Mobile companion app' chapter shows Umoja Pro on Android."),
                size=16, color=INK)
    add_textbox(slide, Inches(0.6), Inches(4.5), Inches(12), Inches(0.5),
                text="Roles covered", size=20, bold=True, color=BRAND)
    role_lines = "\n".join(f"  •  {r.name}  —  {r.who}" for r in ROLES)
    add_textbox(slide, Inches(0.6), Inches(5.1), Inches(12), Inches(2.0),
                text=role_lines, size=13, color=INK)


def chapter_intro(prs: Presentation, ch: Chapter, accent: RGBColor = BRAND,
                  label: str = "ROLE") -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_band(slide, Emu(0), Inches(1.6), accent)
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.5),
                text=label, size=12, bold=True,
                color=RGBColor(0xDD, 0xE3, 0xFF))
    add_textbox(slide, Inches(0.6), Inches(0.7), Inches(12), Inches(0.9),
                text=ch.name, size=32, bold=True,
                color=RGBColor(0xFF, 0xFF, 0xFF))
    add_textbox(slide, Inches(0.6), Inches(1.9), Inches(12), Inches(0.5),
                text="Who does this", size=13, bold=True, color=MUTED)
    add_textbox(slide, Inches(0.6), Inches(2.35), Inches(12), Inches(1.0),
                text=ch.who, size=18, color=INK)
    add_textbox(slide, Inches(0.6), Inches(3.5), Inches(12), Inches(0.5),
                text="Your day, in one paragraph", size=13, bold=True, color=MUTED)
    add_textbox(slide, Inches(0.6), Inches(3.95), Inches(12), Inches(2.5),
                text=ch.daily_summary, size=17, color=INK)
    add_textbox(slide, Inches(0.6), Inches(6.7), Inches(12), Inches(0.5),
                text=f"{len(ch.steps)} steps follow. Tips at the end of the chapter.",
                size=12, color=MUTED)


def step_slide(prs: Presentation, ch: Chapter, idx: int, step: Step) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.4),
                text=f"{ch.name.upper()}   ·   STEP {idx} OF {len(ch.steps)}",
                size=10, bold=True, color=MUTED)
    add_textbox(slide, Inches(0.6), Inches(0.7), Inches(12), Inches(0.8),
                text=step.title, size=24, bold=True, color=INK)
    add_textbox(slide, Inches(0.6), Inches(1.7), Inches(4.8), Inches(5.4),
                text=step.body, size=14, color=INK)
    if step.screen and (SHOTS / step.screen).exists():
        # Mobile screenshots are tall portrait (1080x2400 ≈ 0.45 ratio).
        # Site screenshots are 1366x900 ≈ 1.52.
        if step.screen.startswith("mobile_"):
            pic_left = Inches(7.5)
            pic_top = Inches(1.6)
            pic_h = Inches(5.6)
            slide.shapes.add_picture(str(SHOTS / step.screen), pic_left, pic_top, height=pic_h)
        else:
            pic_left = Inches(5.6)
            pic_top = Inches(1.7)
            pic_w = Inches(7.4)
            slide.shapes.add_picture(str(SHOTS / step.screen), pic_left, pic_top, width=pic_w)
    add_textbox(slide, Inches(0.6), Inches(7.05), Inches(12), Inches(0.3),
                text="Umoja Hardware · User Guide", size=10, color=MUTED)


def tips_slide(prs: Presentation, ch: Chapter) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide, Inches(0.6), Inches(0.3), Inches(12), Inches(0.4),
                text=f"{ch.name.upper()}   ·   TIPS & GOTCHAS",
                size=10, bold=True, color=MUTED)
    add_textbox(slide, Inches(0.6), Inches(0.7), Inches(12), Inches(0.8),
                text="Things that are easy to miss", size=24, bold=True, color=INK)
    tips_text = "\n\n".join(f"⚑   {t}" for t in ch.tips)
    add_textbox(slide, Inches(0.6), Inches(1.8), Inches(12), Inches(5.0),
                text=tips_text, size=15, color=INK)


def closing(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_band(slide, Emu(0), SLIDE_H, INK)
    add_textbox(slide, Inches(0.6), Inches(2.6), Inches(12), Inches(0.9),
                text="Stuck?", size=44, bold=True,
                color=RGBColor(0xFF, 0xFF, 0xFF))
    add_textbox(slide, Inches(0.6), Inches(3.7), Inches(12), Inches(2.5),
                text=("•  Try the menu path written in each step.\n"
                      "•  Refresh the page (Ctrl+R) — the live activity feed reconnects automatically.\n"
                      "•  If a button is missing, your role probably doesn't have permission — ask the System Administrator.\n"
                      "•  Production URL: https://umoja.ehub.co.tz\n"
                      "•  Mobile app: install Umoja Pro from your IT team."),
                size=18, color=RGBColor(0xDD, 0xE3, 0xFF))


def build() -> None:
    prs = Presentation()
    set_size(prs)
    cover(prs)
    intro(prs)
    # End-to-end sale chapter first — it's the most useful demo
    chapter_intro(prs, WORKFLOW, accent=ACCENT, label="END-TO-END WALKTHROUGH")
    for i, s in enumerate(WORKFLOW.steps, start=1):
        step_slide(prs, WORKFLOW, i, s)
    tips_slide(prs, WORKFLOW)
    # Then each role
    for role in ROLES:
        chapter_intro(prs, role, accent=BRAND, label="ROLE")
        for i, s in enumerate(role.steps, start=1):
            step_slide(prs, role, i, s)
        if role.tips:
            tips_slide(prs, role)
    # Mobile companion app
    chapter_intro(prs, MOBILE, accent=RGBColor(0x7A, 0x2A, 0xB8), label="MOBILE APP")
    for i, s in enumerate(MOBILE.steps, start=1):
        step_slide(prs, MOBILE, i, s)
    tips_slide(prs, MOBILE)
    closing(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"wrote {OUT}  ({OUT.stat().st_size // 1024} KB, {len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
